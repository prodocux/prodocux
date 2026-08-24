"""Binary → product-neutral ``prodocux_content_blocks_v1`` (no domain semantics)."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from ..intake.archive import validate_office_archive
from ..intake.docx import MAX_DOCX_BYTES
from ..intake.pdf import MAX_PDF_BYTES, MAX_PDF_PAGES, extract_pdf_bytes
from ..intake.presentation import MAX_PRESENTATION_BYTES
from ..intake.table import MAX_TABLE_BYTES
from ..intake.workbook import MAX_WORKBOOK_BYTES
from .errors import FORMAT_NOT_SUPPORTED, REQUEST_INVALID, RenderContractError
from .validate import validate_content_blocks

MAX_BLOCKS = 200
MAX_ROWS = 500
MAX_COLS = 32
MAX_PARAGRAPHS = 200
MAX_TEXT = 8192
MAX_TEXT_ITEMS = 2000
MAX_SLIDE_PARAS = 50

_FORMAT_LIMITS = {
    "pdf": MAX_PDF_BYTES,
    "csv": MAX_TABLE_BYTES,
    "docx": MAX_DOCX_BYTES,
    "xlsx": MAX_WORKBOOK_BYTES,
    "pptx": MAX_PRESENTATION_BYTES,
}

_SUFFIX_FORMAT = {
    ".pdf": "pdf",
    ".csv": "csv",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".pptx": "pptx",
}

_SHEET_FORBIDDEN = re.compile(r"[\\/?*\[\]:]")


class _Ids:
    def __init__(self) -> None:
        self.n = 0

    def next(self, prefix: str = "b") -> str:
        self.n += 1
        return f"{prefix}{self.n}"


def _clip(text: str, limit: int = MAX_TEXT) -> str:
    return text.replace("\x00", "").strip()[:limit]


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        text = "true" if value else "false"
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif hasattr(value, "isoformat"):
        text = value.isoformat()
    else:
        text = str(value)
    return text.replace("\x00", "")[:MAX_TEXT]


def _sheet_name(raw: str, used: set[str]) -> str:
    name = _SHEET_FORBIDDEN.sub("-", raw).strip() or "Sheet"
    name = name[:64]
    candidate = name
    index = 2
    while candidate in used:
        suffix = f"-{index}"
        candidate = f"{name[: max(1, 64 - len(suffix))]}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def _trim_row(row: list[str]) -> list[str]:
    clipped = row[:MAX_COLS] or [""]
    while len(clipped) > 1 and clipped[-1] == "":
        clipped.pop()
    return clipped or [""]


def flatten_text_items(
    content: Mapping[str, Any], *, limit: int = MAX_TEXT_ITEMS
) -> list[dict[str, str]]:
    """Product-neutral text projection for host adapters (not a content-blocks field)."""
    items: list[dict[str, str]] = []
    for block in content.get("blocks") or []:
        if len(items) >= limit:
            break
        block_id = str(block.get("id") or "b")
        kind = str(block.get("type") or "")
        if kind == "heading":
            items.append(
                {
                    "id": block_id,
                    "type": "heading",
                    "text": str(block.get("text") or ""),
                    "source_locator": f"heading:{block_id}",
                }
            )
        elif kind == "paragraphs":
            for index, paragraph in enumerate(block.get("paragraphs") or [], start=1):
                items.append(
                    {
                        "id": f"{block_id}.{index}",
                        "type": "paragraphs",
                        "text": str(paragraph),
                        "source_locator": f"paragraphs:{block_id}/{index}",
                    }
                )
                if len(items) >= limit:
                    break
        elif kind == "key_values":
            for index, pair in enumerate(block.get("pairs") or [], start=1):
                label = str(pair.get("label") or "")
                value = str(pair.get("value") or "")
                items.append(
                    {
                        "id": f"{block_id}.{index}",
                        "type": "key_values",
                        "text": f"{label}: {value}".strip(),
                        "source_locator": f"key_values:{block_id}/{index}",
                    }
                )
                if len(items) >= limit:
                    break
        elif kind in {"table", "sheet"}:
            name = str(block.get("name") or block_id)
            rows = (block.get("table") or {}).get("rows") or []
            for row_index, row in enumerate(rows, start=1):
                items.append(
                    {
                        "id": f"{block_id}.r{row_index}",
                        "type": f"{kind}_row",
                        "text": " ".join(str(cell) for cell in row).strip(),
                        "source_locator": f"{kind}:{name}/r{row_index}",
                    }
                )
                if len(items) >= limit:
                    break
        elif kind == "slide":
            items.append(
                {
                    "id": block_id,
                    "type": "slide",
                    "text": str(block.get("title") or ""),
                    "source_locator": f"slide:{block_id}/title",
                }
            )
            for index, paragraph in enumerate(block.get("paragraphs") or [], start=1):
                items.append(
                    {
                        "id": f"{block_id}.{index}",
                        "type": "slide",
                        "text": str(paragraph),
                        "source_locator": f"slide:{block_id}/{index}",
                    }
                )
                if len(items) >= limit:
                    break
    return items[:limit]


def extract_content_blocks(filename: str, payload: bytes) -> dict[str, Any]:
    """Parse a 5-format binary into content blocks plus a text projection."""
    name = Path(filename).name
    if name != filename or name in {".", ".."}:
        raise RenderContractError(REQUEST_INVALID, "document_filename must be a plain basename")
    suffix = Path(name).suffix.casefold()
    fmt = _SUFFIX_FORMAT.get(suffix)
    if fmt is None:
        raise RenderContractError(FORMAT_NOT_SUPPORTED, "filename suffix is not a supported extract format")
    if len(payload) > _FORMAT_LIMITS[fmt]:
        raise RenderContractError(REQUEST_INVALID, "document exceeds format byte limit")

    if fmt == "csv":
        content, truncated = _extract_csv(payload, name)
    elif fmt == "xlsx":
        content, truncated = _extract_xlsx(payload)
    elif fmt == "docx":
        content, truncated = _extract_docx(payload)
    elif fmt == "pptx":
        content, truncated = _extract_pptx(payload)
    else:
        content, truncated = _extract_pdf(payload, name)

    title = _clip(Path(name).stem, 512) or "Document"
    content.setdefault("document", {"title": title})
    validate_content_blocks(content)
    from .. import __version__

    return {
        "kernel_version": __version__,
        "format": fmt,
        "source_sha256": hashlib.sha256(payload).hexdigest(),
        "truncated": truncated,
        "content": content,
        "text_items": flatten_text_items(content),
    }


def _extract_csv(payload: bytes, filename: str) -> tuple[dict[str, Any], bool]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise RenderContractError(REQUEST_INVALID, "CSV payload is not UTF-8") from exc
    rows = list(csv.reader(io.StringIO(text)))
    truncated = len(rows) > MAX_ROWS or any(len(row) > MAX_COLS for row in rows)
    table_rows = [_trim_row([_cell(item) for item in row]) for row in rows[:MAX_ROWS]]
    if not table_rows:
        table_rows = [[""]]
    name = _sheet_name(Path(filename).stem or "Sheet", set())
    return {
        "schema_version": "prodocux_content_blocks_v1",
        "blocks": [
            {
                "id": "s1",
                "type": "sheet",
                "name": name,
                "table": {"header_rows": 1 if len(table_rows) > 1 else 0, "rows": table_rows},
            }
        ],
    }, truncated


def _extract_xlsx(payload: bytes) -> tuple[dict[str, Any], bool]:
    from openpyxl import load_workbook

    validate_office_archive(payload, label="XLSX", invalid_message="invalid XLSX workbook")
    truncated = False
    blocks: list[dict[str, Any]] = []
    ids = _Ids()
    used: set[str] = set()
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            if len(blocks) >= MAX_BLOCKS:
                truncated = True
                break
            rows: list[list[str]] = []
            for index, raw_row in enumerate(
                sheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True),
                start=1,
            ):
                if index > MAX_ROWS:
                    truncated = True
                    break
                rows.append(_trim_row([_cell(cell) for cell in raw_row]))
            if sheet.max_row and sheet.max_row > MAX_ROWS:
                truncated = True
            if sheet.max_column and sheet.max_column > MAX_COLS:
                truncated = True
            if not rows or all(all(cell == "" for cell in row) for row in rows):
                continue
            blocks.append(
                {
                    "id": ids.next("s"),
                    "type": "sheet",
                    "name": _sheet_name(str(sheet.title), used),
                    "table": {"header_rows": 1 if len(rows) > 1 else 0, "rows": rows},
                }
            )
    finally:
        workbook.close()
    if not blocks:
        blocks.append(
            {
                "id": "s1",
                "type": "sheet",
                "name": "Sheet",
                "table": {"header_rows": 0, "rows": [[""]]},
            }
        )
    return {"schema_version": "prodocux_content_blocks_v1", "blocks": blocks[:MAX_BLOCKS]}, truncated


def _extract_docx(payload: bytes) -> tuple[dict[str, Any], bool]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    validate_office_archive(payload, label="DOCX", invalid_message="invalid DOCX document")
    document = Document(io.BytesIO(payload))
    ids = _Ids()
    blocks: list[dict[str, Any]] = []
    pending: list[str] = []
    truncated = False

    def flush_paragraphs() -> None:
        if pending and len(blocks) < MAX_BLOCKS:
            blocks.append(
                {
                    "id": ids.next("p"),
                    "type": "paragraphs",
                    "paragraphs": pending[:MAX_PARAGRAPHS],
                }
            )
        pending.clear()

    for child in document.element.body.iterchildren():
        if len(blocks) >= MAX_BLOCKS:
            truncated = True
            break
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = _clip(paragraph.text)
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            if style.casefold().startswith("heading"):
                flush_paragraphs()
                level = 1
                digits = "".join(ch for ch in style if ch.isdigit())
                if digits:
                    level = max(1, min(6, int(digits)))
                blocks.append({"id": ids.next("h"), "type": "heading", "level": level, "text": text})
            elif len(pending) < MAX_PARAGRAPHS:
                pending.append(text)
            else:
                truncated = True
                flush_paragraphs()
                pending.append(text)
        elif child.tag == qn("w:tbl"):
            flush_paragraphs()
            table = Table(child, document)
            rows: list[list[str]] = []
            for row in table.rows[:MAX_ROWS]:
                rows.append(_trim_row([_cell(cell.text) for cell in row.cells[:MAX_COLS]]))
            if len(table.rows) > MAX_ROWS or any(len(row.cells) > MAX_COLS for row in table.rows):
                truncated = True
            if rows:
                blocks.append(
                    {
                        "id": ids.next("t"),
                        "type": "table",
                        "table": {"header_rows": 1 if len(rows) > 1 else 0, "rows": rows},
                    }
                )
    flush_paragraphs()
    if not blocks:
        blocks.append({"id": "h1", "type": "heading", "level": 1, "text": "Document"})
    return {"schema_version": "prodocux_content_blocks_v1", "blocks": blocks[:MAX_BLOCKS]}, truncated


def _extract_pptx(payload: bytes) -> tuple[dict[str, Any], bool]:
    from pptx import Presentation

    validate_office_archive(payload, label="PPTX", invalid_message="invalid PPTX presentation")
    presentation = Presentation(io.BytesIO(payload))
    ids = _Ids()
    blocks: list[dict[str, Any]] = []
    truncated = False
    for index, slide in enumerate(presentation.slides, start=1):
        if len(blocks) >= MAX_BLOCKS:
            truncated = True
            break
        title_shape = slide.shapes.title
        title = _clip(title_shape.text) if title_shape is not None else ""
        paragraphs: list[str] = []
        for shape in slide.shapes:
            if title_shape is not None and shape == title_shape:
                continue
            if getattr(shape, "has_text_frame", False):
                text = _clip(shape.text)
                if text and len(paragraphs) < MAX_SLIDE_PARAS:
                    paragraphs.append(text)
            if getattr(shape, "has_table", False) and len(blocks) < MAX_BLOCKS - 1:
                rows: list[list[str]] = []
                for row in shape.table.rows[:MAX_ROWS]:
                    rows.append(_trim_row([_cell(cell.text) for cell in row.cells[:MAX_COLS]]))
                if rows:
                    # Keep tables as sibling blocks; slide itself stays title + paragraphs.
                    blocks.append(
                        {
                            "id": ids.next("t"),
                            "type": "table",
                            "table": {"header_rows": 1 if len(rows) > 1 else 0, "rows": rows},
                        }
                    )
        blocks.append(
            {
                "id": ids.next("d"),
                "type": "slide",
                "title": title or f"Slide {index}",
                **({"paragraphs": paragraphs} if paragraphs else {}),
            }
        )
    if not blocks:
        blocks.append({"id": "d1", "type": "slide", "title": "Presentation"})
    return {"schema_version": "prodocux_content_blocks_v1", "blocks": blocks[:MAX_BLOCKS]}, truncated


def _extract_pdf(payload: bytes, filename: str) -> tuple[dict[str, Any], bool]:
    pages, truncated = extract_pdf_bytes(
        payload, filename=filename, max_pages=MAX_PDF_PAGES
    )
    ids = _Ids()
    blocks: list[dict[str, Any]] = []
    for page in pages:
        if len(blocks) >= MAX_BLOCKS:
            truncated = True
            break
        text = _clip(str(page.get("text") or ""))
        if not text:
            continue
        paragraphs = [part[:MAX_TEXT] for part in text.splitlines() if part.strip()][:MAX_PARAGRAPHS]
        if not paragraphs:
            continue
        blocks.append({"id": ids.next("p"), "type": "paragraphs", "paragraphs": paragraphs})
    if not blocks:
        title = _clip(Path(filename).stem, 512) or "Document"
        blocks.append({"id": "h1", "type": "heading", "level": 1, "text": title})
    return {"schema_version": "prodocux_content_blocks_v1", "blocks": blocks[:MAX_BLOCKS]}, truncated
