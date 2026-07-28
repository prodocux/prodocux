"""#7 Number-consistency audit engine (deterministic, no LLM calls).

Checks:
- table_total_mismatch: table "合計/小計/總計" (total/subtotal) row != sum of the
  column above it
- amount_words_mismatch: amount in Chinese numeral words != the Arabic-numeral
  amount
- date_format_inconsistent: the document mixes multiple date formats
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field, asdict
from typing import Any, Dict, List, Optional

from docx import Document


@dataclass
class Finding:
    check: str                  # message code, also used as the check identifier (localized by the presentation layer)
    severity: str               # high | medium | low
    loc: str                    # location code (e.g. loc.table_col / loc.text / loc.document)
    loc_params: Dict[str, Any] = field(default_factory=dict)
    params: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


# --------------------------------------------------------------------------
# Parsing Chinese numeral-word amounts
# --------------------------------------------------------------------------
_CN_DIGIT = {
    "零": 0, "〇": 0, "一": 1, "壹": 1, "二": 2, "貳": 2, "两": 2, "兩": 2,
    "三": 3, "參": 3, "叁": 3, "四": 4, "肆": 4, "五": 5, "伍": 5,
    "六": 6, "陸": 6, "七": 7, "柒": 7, "八": 8, "捌": 8, "九": 9, "玖": 9,
}
_CN_UNIT = {"十": 10, "拾": 10, "百": 100, "佰": 100, "千": 1000, "仟": 1000}
_CN_BIG = {"萬": 10000, "万": 10000, "億": 10 ** 8, "亿": 10 ** 8}

_CN_AMOUNT_RE = re.compile(
    r"[零〇一壹二貳兩两三參叁四肆五伍六陸七柒八捌九玖十拾百佰千仟萬万億亿]{2,}"
)
_CURRENCY_HINT = re.compile(r"(元|圓|NT\$?|新台幣|新臺幣|金額|價|＄|\$)")
_ARABIC_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")


def parse_cn_number(s: str) -> int:
    total = 0
    section = 0
    number = 0
    for ch in s:
        if ch in _CN_DIGIT:
            number = _CN_DIGIT[ch]
        elif ch in _CN_UNIT:
            unit = _CN_UNIT[ch]
            if number == 0:
                number = 1  # "十" (ten) = 10
            section += number * unit
            number = 0
        elif ch in _CN_BIG:
            section += number
            total += section * _CN_BIG[ch]
            section = 0
            number = 0
    return total + section + number


def _to_number(s: str) -> Optional[float]:
    if s is None:
        return None
    s = s.strip().replace(",", "").replace("，", "")
    s = re.sub(r"[^\d.\-]", "", s)
    if s in ("", "-", ".", "--"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


# --------------------------------------------------------------------------
# Checks
# --------------------------------------------------------------------------
_TOTAL_KEYWORDS = ("合計", "小計", "總計", "總額", "合 計", "total", "sum", "subtotal")


def check_table_totals(doc: Document) -> List[Finding]:
    findings: List[Finding] = []
    for ti, table in enumerate(doc.tables):
        grid = [[c.text for c in row.cells] for row in table.rows]
        if len(grid) < 2:
            continue
        for ri, row in enumerate(grid):
            label = (row[0] or "").strip().lower()
            if not any(k in label for k in _TOTAL_KEYWORDS):
                continue
            for ci in range(1, len(row)):
                total_val = _to_number(row[ci])
                if total_val is None:
                    continue
                nums = [
                    _to_number(grid[r][ci])
                    for r in range(ri)
                    if ci < len(grid[r])
                ]
                nums = [v for v in nums if v is not None]
                if len(nums) < 2:
                    continue
                s = sum(nums)
                if abs(s - total_val) > 0.01 + abs(total_val) * 1e-9:
                    findings.append(Finding(
                        "table_total_mismatch", "high",
                        "loc.table_col", {"table": ti + 1, "col": ci + 1},
                        {"total": total_val, "sum": s, "diff": total_val - s},
                    ))
            break  # only take the first total row per table
    return findings


def check_amount_in_words(doc: Document) -> List[Finding]:
    findings: List[Finding] = []
    for p in doc.paragraphs:
        text = p.text
        if not _CURRENCY_HINT.search(text):
            continue
        cn_matches = list(_CN_AMOUNT_RE.finditer(text))
        ar_matches = list(_ARABIC_RE.finditer(text))
        if not cn_matches or not ar_matches:
            continue
        for cm in cn_matches:
            cn_val = parse_cn_number(cm.group())
            if cn_val == 0:
                continue
            nearest = min(ar_matches, key=lambda m: abs(m.start() - cm.start()))
            ar_val = _to_number(nearest.group())
            if ar_val is None:
                continue
            if abs(cn_val - ar_val) > 0.5:
                findings.append(Finding(
                    "amount_words_mismatch", "high",
                    "loc.text", {"snippet": text.strip()[:40]},
                    {"words": cm.group(), "words_value": cn_val, "number_value": ar_val},
                ))
    return findings


_DATE_PATTERNS = {
    "YYYY/MM/DD": re.compile(r"\b\d{4}/\d{1,2}/\d{1,2}\b"),
    "YYYY-MM-DD": re.compile(r"\b\d{4}-\d{1,2}-\d{1,2}\b"),
    "YYYY年MM月DD日": re.compile(r"\d{4}年\d{1,2}月\d{1,2}日"),
    "MM/DD/YYYY": re.compile(r"\b\d{1,2}/\d{1,2}/\d{4}\b"),
}


def check_date_formats(doc: Document) -> List[Finding]:
    text = "\n".join(p.text for p in doc.paragraphs)
    styles = sorted(name for name, pat in _DATE_PATTERNS.items() if pat.search(text))
    if len(styles) > 1:
        return [Finding(
            "date_format_inconsistent", "medium",
            "loc.document", {}, {"styles": ", ".join(styles)},
        )]
    return []


def audit_numbers(document_path: str) -> Dict:
    doc = Document(document_path)
    findings: List[Finding] = []
    findings += check_table_totals(doc)
    findings += check_amount_in_words(doc)
    findings += check_date_formats(doc)
    return {
        "file": document_path,
        "passed": len(findings) == 0,
        "findings": [f.to_dict() for f in findings],
    }
