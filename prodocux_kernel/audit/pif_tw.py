"""#9 Taiwan PIF compliance audit engine (deterministic, no LLM calls).

Based on the 16 required data items in Article 3 of the MOHW's "Regulations
Governing the Management of Cosmetic Product Information Files", checks
against the checklist and template mapping:
- whether all 16 sections exist and have substantive content
- placeholder / pending-confirmation markers
- section keyword coverage (heuristic)
- sub-requirements of item 16 (safety assessment conclusion + signer
  qualifications)
- L0 structural checks (TOC field, blank pages, cross-references)
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Union

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from ..docops.invariants import validate_structure
from ..templating.mapping import MappingConfig
from ..templating.profile import TemplateProfile, extract_profile

PathLike = Union[str, Path]

_STRUCTURE_SEVERITY = {
    "valid_docx": "critical",
    "toc_is_field": "high",
    "cross_references_valid": "high",
    "no_blank_pages": "medium",
}


@dataclass
class PifFinding:
    check: str
    severity: str
    section_id: str = ""
    loc: str = "loc.document"
    loc_params: Dict[str, Any] = field(default_factory=dict)
    params: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


def load_checklist(path: PathLike) -> Dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _cell_text(cell) -> str:
    return " ".join(p.text for p in cell.paragraphs if p.text.strip())


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        rows.append(" | ".join(_cell_text(c) for c in row.cells))
    return "\n".join(rows)


def extract_section_bodies(doc: Document, profile: TemplateProfile) -> Dict[str, str]:
    """Extract the paragraph and table text beneath each section heading."""
    children = list(doc.element.body.iterchildren())
    bodies: Dict[str, str] = {}
    for slot in profile.sections:
        if not slot.found or slot.para_start is None:
            continue
        end = slot.para_end if slot.para_end is not None else len(children)
        parts: List[str] = []
        for child in children[slot.para_start + 1: end]:
            if child.tag == qn("w:p"):
                text = Paragraph(child, doc).text.strip()
                if text:
                    parts.append(text)
            elif child.tag == qn("w:tbl"):
                from docx.table import Table
                parts.append(_table_text(Table(child, doc)))
        bodies[slot.section_id] = "\n".join(parts).strip()
    return bodies


def _is_placeholder(text: str, patterns: Sequence[str]) -> bool:
    if not text.strip():
        return True
    lower = text.lower()
    for pat in patterns:
        if pat.lower() in lower:
            return True
    if len(text.strip()) < 8:
        return True
    return False


def _has_keyword(text: str, keywords: Sequence[str]) -> bool:
    lower = text.lower()
    return any(kw.lower() in lower for kw in keywords)


def _keyword_hits(text: str, keywords: Sequence[str]) -> List[str]:
    lower = text.lower()
    return [kw for kw in keywords if kw.lower() in lower]


def audit_pif_tw(
    document_path: PathLike,
    mapping: MappingConfig,
    checklist: Dict[str, Any],
) -> Dict[str, Any]:
    """Run the Taiwan PIF compliance audit and return a language-neutral structured result."""
    doc = Document(str(document_path))
    profile = extract_profile(
        doc,
        mapping.section_specs(),
        mapping.heading_styles,
        mapping.body_style,
    )
    bodies = extract_section_bodies(doc, profile)
    placeholders = checklist.get(
        "empty_placeholder_patterns",
        checklist.get("placeholder_patterns", []),
    )
    review_patterns = checklist.get("review_patterns", [])
    required = {s["id"]: s for s in checklist.get("required_sections", [])}

    findings: List[PifFinding] = []

    # --- Presence and content of the 16 sections ---
    for spec in checklist.get("required_sections", []):
        sid = spec["id"]
        slot = profile.section(sid)
        title = spec.get("title", sid)

        if slot is None or not slot.found:
            findings.append(PifFinding(
                check="section_missing",
                severity="critical",
                section_id=sid,
                loc="loc.section",
                loc_params={"id": sid, "title": title},
                params={"title": title},
            ))
            continue

        body = bodies.get(sid, "")
        if _is_placeholder(body, placeholders):
            findings.append(PifFinding(
                check="section_empty_or_placeholder",
                severity="high",
                section_id=sid,
                loc="loc.section",
                loc_params={"id": sid, "title": title},
                params={"title": title, "chars": len(body)},
            ))
            continue

        keywords = spec.get("keywords", [])
        if keywords and not _has_keyword(body, keywords):
            findings.append(PifFinding(
                check="section_keywords_missing",
                severity="medium",
                section_id=sid,
                loc="loc.section",
                loc_params={"id": sid, "title": title},
                params={"title": title, "sample": keywords[:3]},
            ))

        for req_kw in spec.get("required_keywords", []):
            if req_kw.lower() not in body.lower():
                findings.append(PifFinding(
                    check="section_required_keyword_missing",
                    severity="high",
                    section_id=sid,
                    loc="loc.section",
                    loc_params={"id": sid, "title": title},
                    params={"title": title, "keyword": req_kw},
                ))

        for pat in review_patterns:
            if pat in body:
                findings.append(PifFinding(
                    check="review_marker_in_section",
                    severity="medium",
                    section_id=sid,
                    loc="loc.section",
                    loc_params={"id": sid, "title": title},
                    params={"title": title, "marker": pat},
                ))
                break

    # --- Sub-requirements of item 16 ---
    safety_body = bodies.get("16_safety", "")
    if safety_body and not _is_placeholder(safety_body, placeholders):
        for sub in checklist.get("section_16_subrequirements", []):
            if not _has_keyword(safety_body, sub.get("keywords", [])):
                findings.append(PifFinding(
                    check="safety_subrequirement_missing",
                    severity="high",
                    section_id="16_safety",
                    loc="loc.section",
                    loc_params={"id": sub["id"], "title": "16. 產品安全資料"},
                    params={"sub_id": sub["id"]},
                ))

    # --- L0 structural checks ---
    for inv in validate_structure(str(document_path)):
        if inv.id not in checklist.get("structure_checks", []):
            continue
        if inv.status != "checked" or inv.passed:
            continue
        findings.append(PifFinding(
            check=f"structure.{inv.id}",
            severity=_STRUCTURE_SEVERITY.get(inv.id, "medium"),
            section_id="",
            loc="loc.document",
            params={"code": inv.code, **inv.params},
        ))

    failed = [f for f in findings]
    max_sev = "none"
    rank = {"critical": 4, "high": 3, "medium": 2, "low": 1, "none": 0}
    if failed:
        max_sev = max(failed, key=lambda f: rank[f.severity]).severity

    present = sum(1 for s in profile.sections if s.found)
    return {
        "file": str(document_path),
        "passed": len(failed) == 0,
        "sections_found": present,
        "sections_required": len(required),
        "findings": [f.to_dict() for f in findings],
        "fail_count": len(failed),
        "max_severity": max_sev,
        "regulation": checklist.get("source", {}),
        "section_bodies": {k: len(v) for k, v in bodies.items()},
    }
