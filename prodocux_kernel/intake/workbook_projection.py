"""Bounded deterministic XLSX worksheet-row projection."""

from __future__ import annotations

import hashlib
import io
import zipfile
from collections.abc import Mapping
from copy import deepcopy
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from prodocux_kernel.intake.archive import validate_office_archive
from prodocux_kernel.intake.errors import SourceTooLargeError
from prodocux_kernel.intake.projection_validate import validate_continuable_projection
from prodocux_kernel.intake.workbook import MAX_WORKBOOK_BYTES, _cell_value

MAX_XLSX_ROWS_PER_RANGE = 500
MAX_XLSX_COLUMNS = 32
PARSER_CONTRACT_NAME = "prodocux_xlsx_worksheet_row_projection_v1"
PARSER_CONTRACT_VERSION = "1"


def extract_xlsx_continuable_projection(
    payload: bytes,
    *,
    cursor: Mapping[str, Any] | None = None,
    max_rows: int = MAX_XLSX_ROWS_PER_RANGE,
) -> dict[str, Any]:
    if not payload:
        raise ValueError("document payload is empty")
    if len(payload) > MAX_WORKBOOK_BYTES:
        raise SourceTooLargeError(f"XLSX exceeds {MAX_WORKBOOK_BYTES} bytes")
    if not 1 <= max_rows <= MAX_XLSX_ROWS_PER_RANGE:
        raise ValueError("max_rows must be between 1 and 500")
    validate_office_archive(
        payload, label="XLSX", invalid_message="invalid XLSX workbook"
    )
    source_sha256 = hashlib.sha256(payload).hexdigest()
    try:
        formulas = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
        values = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise ValueError("invalid XLSX workbook") from exc
    try:
        sheet_index = 0
        start_row = 1
        if cursor is not None:
            expected = {
                "schema_version": "prodocux_projection_cursor_v1",
                "source_sha256": source_sha256,
                "parser_contract_name": PARSER_CONTRACT_NAME,
                "parser_contract_version": PARSER_CONTRACT_VERSION,
            }
            for name, value in expected.items():
                if cursor.get(name) != value:
                    raise ValueError(f"cursor {name} does not match source projection")
            sheet_index = cursor.get("next_sheet_index")
            start_row = cursor.get("next_row")
            if (
                not isinstance(sheet_index, int)
                or isinstance(sheet_index, bool)
                or sheet_index < 0
                or not isinstance(start_row, int)
                or isinstance(start_row, bool)
                or start_row < 1
            ):
                raise ValueError("cursor worksheet range is invalid")
        if sheet_index >= len(formulas.worksheets):
            raise ValueError("cursor sheet index exceeds workbook sheet count")
        formula_sheet = formulas.worksheets[sheet_index]
        value_sheet = values[formula_sheet.title]
        if cursor is not None and cursor.get("next_sheet_name") != formula_sheet.title:
            raise ValueError("cursor sheet name does not match workbook")
        max_row = formula_sheet.max_row
        if start_row > max_row:
            raise ValueError("cursor next_row exceeds worksheet row count")
        end_row = min(start_row + max_rows, max_row + 1)
        column_count = min(formula_sheet.max_column, MAX_XLSX_COLUMNS)
        omissions = (
            ["columns_beyond_limit"]
            if formula_sheet.max_column > MAX_XLSX_COLUMNS
            else []
        )
        rows: list[list[Any]] = []
        formula_cells: list[dict[str, Any]] = []
        for row_index, row in enumerate(
            formula_sheet.iter_rows(
                min_row=start_row,
                max_row=end_row - 1,
                max_col=column_count,
            ),
            start=start_row,
        ):
            projected: list[Any] = []
            for column_index, cell in enumerate(row, start=1):
                cached = value_sheet.cell(row=row_index, column=column_index).value
                projected.append(
                    _cell_value(cached if cached is not None else cell.value)
                )
                if getattr(cell, "data_type", None) == "f":
                    formula_cells.append(
                        {
                            "cell": f"{get_column_letter(column_index)}{row_index}",
                            "formula": str(cell.value),
                            "cached_value": _cell_value(cached),
                        }
                    )
            rows.append(projected)
        next_cursor = None
        if end_row <= max_row:
            next_cursor = _cursor(
                source_sha256, sheet_index, formula_sheet.title, end_row
            )
        elif sheet_index + 1 < len(formulas.worksheets):
            next_sheet = formulas.worksheets[sheet_index + 1]
            next_cursor = _cursor(source_sha256, sheet_index + 1, next_sheet.title, 1)
        disposition = (
            "partial_unknown"
            if omissions
            else "partial_known"
            if next_cursor is not None
            else "complete"
        )
        result = {
            "schema_version": "prodocux_xlsx_continuable_projection_v1",
            "source_sha256": source_sha256,
            "parser_contract": {
                "name": PARSER_CONTRACT_NAME,
                "version": PARSER_CONTRACT_VERSION,
            },
            "range": {
                "unit": "worksheet_row",
                "sheet_index": sheet_index,
                "sheet_name": formula_sheet.title,
                "start": start_row,
                "end": end_row,
                "requested_max_rows": max_rows,
            },
            "rows": rows,
            "formula_cells": formula_cells,
            "next_cursor": deepcopy(next_cursor),
            "coverage": {
                "disposition": disposition,
                "known_total_sheets": len(formulas.worksheets),
                "known_sheet_rows": max_row,
                "continuation_available": next_cursor is not None,
                "omitted_content_classes": omissions,
            },
            "counts": {
                "returned_rows": len(rows),
                "returned_columns": column_count,
            },
        }
        validate_continuable_projection(result)
        return result
    finally:
        formulas.close()
        values.close()


def _cursor(
    source_sha256: str, sheet_index: int, sheet_name: str, next_row: int
) -> dict[str, Any]:
    return {
        "schema_version": "prodocux_projection_cursor_v1",
        "source_sha256": source_sha256,
        "parser_contract_name": PARSER_CONTRACT_NAME,
        "parser_contract_version": PARSER_CONTRACT_VERSION,
        "next_sheet_index": sheet_index,
        "next_sheet_name": sheet_name,
        "next_row": next_row,
    }
