"""Bounded deterministic PPTX slide-range projection."""

from __future__ import annotations

import hashlib
import io
import zipfile
from collections.abc import Mapping
from copy import deepcopy
from itertools import islice
from typing import Any

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from prodocux_kernel.intake.archive import validate_office_archive
from prodocux_kernel.intake.errors import SourceTooLargeError
from prodocux_kernel.intake.presentation import (
    MAX_PRESENTATION_BYTES,
    MAX_SHAPES_PER_SLIDE,
    MAX_TABLE_COLUMNS,
    MAX_TABLE_ROWS,
    _clean,
)
from prodocux_kernel.intake.projection_validate import validate_continuable_projection

MAX_PPTX_SLIDES_PER_RANGE = 50
PARSER_CONTRACT_NAME = "prodocux_pptx_slide_projection_v1"
PARSER_CONTRACT_VERSION = "1"


def extract_pptx_continuable_projection(
    payload: bytes,
    *,
    cursor: Mapping[str, Any] | None = None,
    max_slides: int = MAX_PPTX_SLIDES_PER_RANGE,
) -> dict[str, Any]:
    if not payload:
        raise ValueError("document payload is empty")
    if len(payload) > MAX_PRESENTATION_BYTES:
        raise SourceTooLargeError(f"PPTX exceeds {MAX_PRESENTATION_BYTES} bytes")
    if not 1 <= max_slides <= MAX_PPTX_SLIDES_PER_RANGE:
        raise ValueError("max_slides must be between 1 and 50")
    validate_office_archive(
        payload, label="PPTX", invalid_message="invalid PPTX presentation"
    )
    source_sha256 = hashlib.sha256(payload).hexdigest()
    try:
        presentation = Presentation(io.BytesIO(payload))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError) as exc:
        raise ValueError("invalid PPTX presentation") from exc
    total_slides = len(presentation.slides)
    if total_slides == 0:
        raise ValueError("PPTX contains no slides")
    start_slide = 1
    if cursor is not None:
        expected = {
            "schema_version": "prodocux_projection_cursor_v1",
            "source_sha256": source_sha256,
            "parser_contract_name": PARSER_CONTRACT_NAME,
            "parser_contract_version": PARSER_CONTRACT_VERSION,
        }
        for name, value in expected.items():
            if cursor.get(name) != value:
                raise ValueError(f"cursor {name} does not match source projection")
        start_slide = cursor.get("next_slide")
        if (
            not isinstance(start_slide, int)
            or isinstance(start_slide, bool)
            or start_slide < 1
        ):
            raise ValueError("cursor next_slide must be a positive integer")
    if start_slide > total_slides:
        raise ValueError("cursor next_slide exceeds presentation slide count")
    end_slide = min(start_slide + max_slides, total_slides + 1)
    slides: list[dict[str, Any]] = []
    omissions: set[str] = set()
    for slide_number in range(start_slide, end_slide):
        slides.append(
            _project_slide(
                presentation.slides[slide_number - 1], slide_number, omissions
            )
        )
    next_cursor = None
    if end_slide <= total_slides:
        next_cursor = {
            "schema_version": "prodocux_projection_cursor_v1",
            "source_sha256": source_sha256,
            "parser_contract_name": PARSER_CONTRACT_NAME,
            "parser_contract_version": PARSER_CONTRACT_VERSION,
            "next_slide": end_slide,
        }
    disposition = (
        "partial_unknown"
        if omissions
        else "partial_known"
        if next_cursor is not None
        else "complete"
    )
    result = {
        "schema_version": "prodocux_pptx_continuable_projection_v1",
        "source_sha256": source_sha256,
        "parser_contract": {
            "name": PARSER_CONTRACT_NAME,
            "version": PARSER_CONTRACT_VERSION,
        },
        "range": {
            "unit": "slide",
            "start": start_slide,
            "end": end_slide,
            "requested_max_slides": max_slides,
        },
        "slides": slides,
        "next_cursor": deepcopy(next_cursor),
        "coverage": {
            "disposition": disposition,
            "known_total_slides": total_slides,
            "continuation_available": next_cursor is not None,
            "omitted_content_classes": sorted(omissions),
        },
        "counts": {"returned_slides": len(slides)},
    }
    validate_continuable_projection(result)
    return result


def _project_slide(
    slide: Any, slide_number: int, omissions: set[str]
) -> dict[str, Any]:
    shapes: list[dict[str, Any]] = []
    window = list(islice(slide.shapes, MAX_SHAPES_PER_SLIDE + 1))
    if len(window) > MAX_SHAPES_PER_SLIDE:
        omissions.add("shapes_beyond_limit")
    for shape in window[:MAX_SHAPES_PER_SLIDE]:
        item: dict[str, Any] = {
            "shape_id": shape.shape_id,
            "shape_name": shape.name,
            "text": _clean(shape.text)
            if getattr(shape, "has_text_frame", False)
            else "",
            "is_picture": getattr(shape, "shape_type", None) == 13,
            "table": None,
        }
        if getattr(shape, "has_table", False):
            table = shape.table
            if (
                len(table.rows) > MAX_TABLE_ROWS
                or len(table.columns) > MAX_TABLE_COLUMNS
            ):
                omissions.add("table_cells_beyond_limit")
            item["table"] = [
                [_clean(cell.text) for cell in islice(row.cells, MAX_TABLE_COLUMNS)]
                for row in islice(table.rows, MAX_TABLE_ROWS)
            ]
        shapes.append(item)
    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = _clean(slide.notes_slide.notes_text_frame.text)
    return {"slide_number": slide_number, "shapes": shapes, "speaker_notes": notes}
