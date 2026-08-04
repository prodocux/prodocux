"""Deterministic XLSX workbook profiling (no schedule interpretation)."""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter

from .archive import validate_office_archive

WORKBOOK_PROFILE_SCHEMA = "prodocux_workbook_profile_v1"
MAX_WORKBOOK_BYTES = 16 * 1024 * 1024
MAX_PREVIEW_ROWS = 20
MAX_PREVIEW_COLUMNS = 30


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def profile_xlsx_bytes(raw: bytes, *, filename: str) -> dict[str, Any]:
    if len(raw) > MAX_WORKBOOK_BYTES:
        raise ValueError(f"XLSX exceeds {MAX_WORKBOOK_BYTES} bytes")
    validate_office_archive(raw, label="XLSX", invalid_message="invalid XLSX workbook")
    from io import BytesIO

    try:
        formulas = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        values = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise ValueError("invalid XLSX workbook") from exc
    sheets: list[dict[str, Any]] = []
    try:
        for formula_sheet in formulas.worksheets:
            value_sheet = values[formula_sheet.title]
            preview: list[list[Any]] = []
            formula_cells: list[dict[str, Any]] = []
            for row_index, row in enumerate(
                formula_sheet.iter_rows(
                    min_row=1,
                    max_row=min(formula_sheet.max_row, MAX_PREVIEW_ROWS),
                    max_col=min(formula_sheet.max_column, MAX_PREVIEW_COLUMNS),
                ),
                start=1,
            ):
                preview_row: list[Any] = []
                for column_index, cell in enumerate(row, start=1):
                    coordinate = f"{get_column_letter(column_index)}{row_index}"
                    display_value = value_sheet.cell(
                        row=row_index, column=column_index
                    ).value
                    preview_row.append(_cell_value(display_value if display_value is not None else cell.value))
                    if getattr(cell, "data_type", None) == "f":
                        formula_cells.append(
                            {
                                "cell": coordinate,
                                "formula": str(cell.value),
                                "cached_value": _cell_value(display_value),
                            }
                        )
                preview.append(preview_row)
            sheets.append(
                {
                    "name": formula_sheet.title,
                    "max_row": formula_sheet.max_row,
                    "max_column": formula_sheet.max_column,
                    "preview": preview,
                    "formula_cells": formula_cells,
                    "preview_truncated": (
                        formula_sheet.max_row > MAX_PREVIEW_ROWS
                        or formula_sheet.max_column > MAX_PREVIEW_COLUMNS
                    ),
                }
            )
    finally:
        formulas.close()
        values.close()
    return {
        "schema_version": WORKBOOK_PROFILE_SCHEMA,
        "source": {
            "name": Path(filename).name,
            "sha256": hashlib.sha256(raw).hexdigest(),
        },
        "sheet_count": len(sheets),
        "sheets": sheets,
        "interpretation": "none",
    }


def profile_xlsx(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    return profile_xlsx_bytes(source.read_bytes(), filename=source.name)
