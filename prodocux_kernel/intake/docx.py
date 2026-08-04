"""Deterministic DOCX content profiling (no semantic interpretation)."""

from __future__ import annotations

import hashlib
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from .archive import validate_office_archive

DOCX_PROFILE_SCHEMA = "prodocux_docx_profile_v1"
MAX_DOCX_BYTES = 16 * 1024 * 1024
MAX_PARAGRAPHS = 500
MAX_TABLES = 50
MAX_TABLE_ROWS = 50
MAX_TABLE_COLUMNS = 30


def _text(value: str) -> str:
    return value.replace("\x00", "").strip()


def profile_docx_bytes(raw: bytes, *, filename: str) -> dict[str, Any]:
    if len(raw) > MAX_DOCX_BYTES:
        raise ValueError(f"DOCX exceeds {MAX_DOCX_BYTES} bytes")
    validate_office_archive(raw, label="DOCX", invalid_message="invalid DOCX document")
    try:
        document = Document(BytesIO(raw))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError) as exc:
        raise ValueError("invalid DOCX document") from exc

    paragraphs: list[dict[str, Any]] = []
    heading_count = 0
    for index, paragraph in enumerate(document.paragraphs[:MAX_PARAGRAPHS], start=1):
        text = _text(paragraph.text)
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.casefold().startswith("heading"):
            heading_count += 1
        if text:
            paragraphs.append({"index": index, "style": style, "text": text})

    tables: list[dict[str, Any]] = []
    for table_index, table in enumerate(document.tables[:MAX_TABLES], start=1):
        rows: list[list[str]] = []
        for row in table.rows[:MAX_TABLE_ROWS]:
            rows.append([_text(cell.text) for cell in row.cells[:MAX_TABLE_COLUMNS]])
        tables.append(
            {
                "index": table_index,
                "row_count": len(table.rows),
                "column_count": max((len(row.cells) for row in table.rows), default=0),
                "preview": rows,
                "preview_truncated": (
                    len(table.rows) > MAX_TABLE_ROWS
                    or any(len(row.cells) > MAX_TABLE_COLUMNS for row in table.rows)
                ),
            }
        )

    headers: list[str] = []
    footers: list[str] = []
    for section in document.sections:
        header = "\n".join(_text(item.text) for item in section.header.paragraphs if _text(item.text))
        footer = "\n".join(_text(item.text) for item in section.footer.paragraphs if _text(item.text))
        if header and header not in headers:
            headers.append(header)
        if footer and footer not in footers:
            footers.append(footer)

    return {
        "schema_version": DOCX_PROFILE_SCHEMA,
        "source": {
            "name": Path(filename).name,
            "sha256": hashlib.sha256(raw).hexdigest(),
        },
        "paragraph_count": len(document.paragraphs),
        "heading_count": heading_count,
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "paragraphs": paragraphs,
        "tables": tables,
        "headers": headers,
        "footers": footers,
        "preview_truncated": (
            len(document.paragraphs) > MAX_PARAGRAPHS
            or len(document.tables) > MAX_TABLES
        ),
        "interpretation": "none",
    }


def profile_docx(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    return profile_docx_bytes(source.read_bytes(), filename=source.name)
