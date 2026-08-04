"""Deterministic PPTX presentation profiling (no semantic interpretation)."""

from __future__ import annotations

import hashlib
from itertools import islice
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from .archive import validate_office_archive

PRESENTATION_PROFILE_SCHEMA = "prodocux_presentation_profile_v1"
MAX_PRESENTATION_BYTES = 32 * 1024 * 1024
MAX_SLIDES = 300
MAX_SHAPES_PER_SLIDE = 300
MAX_TABLE_ROWS = 50
MAX_TABLE_COLUMNS = 30


def _clean(value: str) -> str:
    return value.replace("\x00", "").strip()


def profile_pptx_bytes(raw: bytes, *, filename: str) -> dict[str, Any]:
    if len(raw) > MAX_PRESENTATION_BYTES:
        raise ValueError(f"PPTX exceeds {MAX_PRESENTATION_BYTES} bytes")
    validate_office_archive(
        raw, label="PPTX", invalid_message="invalid PPTX presentation"
    )
    try:
        presentation = Presentation(BytesIO(raw))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError) as exc:
        raise ValueError("invalid PPTX presentation") from exc

    slides: list[dict[str, Any]] = []
    for slide_number, slide in enumerate(
        islice(presentation.slides, MAX_SLIDES), start=1
    ):
        title_shape = slide.shapes.title
        title = _clean(title_shape.text) if title_shape is not None else ""
        texts: list[str] = []
        tables: list[dict[str, Any]] = []
        image_count = 0
        shape_window = list(islice(slide.shapes, MAX_SHAPES_PER_SLIDE + 1))
        shapes_truncated = len(shape_window) > MAX_SHAPES_PER_SLIDE
        for shape in shape_window[:MAX_SHAPES_PER_SLIDE]:
            if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                image_count += 1
            if getattr(shape, "has_text_frame", False):
                text = _clean(shape.text)
                if text and text != title:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                table = shape.table
                preview = [
                    [_clean(cell.text) for cell in islice(row.cells, MAX_TABLE_COLUMNS)]
                    for row in islice(table.rows, MAX_TABLE_ROWS)
                ]
                tables.append(
                    {
                        "row_count": len(table.rows),
                        "column_count": len(table.columns),
                        "preview": preview,
                        "preview_truncated": (
                            len(table.rows) > MAX_TABLE_ROWS
                            or len(table.columns) > MAX_TABLE_COLUMNS
                        ),
                    }
                )
        notes = ""
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes = _clean(notes_frame.text)
        slides.append(
            {
                "slide_number": slide_number,
                "title": title,
                "texts": texts,
                "speaker_notes": notes,
                "tables": tables,
                "image_count": image_count,
                "shape_count": (
                    MAX_SHAPES_PER_SLIDE + 1
                    if shapes_truncated
                    else len(shape_window)
                ),
                "shape_count_is_lower_bound": shapes_truncated,
                "preview_truncated": shapes_truncated,
            }
        )
    return {
        "schema_version": PRESENTATION_PROFILE_SCHEMA,
        "source": {
            "name": Path(filename).name,
            "sha256": hashlib.sha256(raw).hexdigest(),
        },
        "slide_count": len(presentation.slides),
        "slides": slides,
        "preview_truncated": len(presentation.slides) > MAX_SLIDES,
        "interpretation": "none",
    }


def profile_pptx(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    return profile_pptx_bytes(source.read_bytes(), filename=source.name)
