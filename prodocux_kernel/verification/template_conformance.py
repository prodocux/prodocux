"""Deterministic DOCX table profiling and template conformance."""

from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from collections.abc import Mapping
from importlib import resources
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jsonschema import Draft202012Validator
from openpyxl import load_workbook
from pptx import Presentation
from referencing import Registry, Resource

from prodocux_kernel.intake.archive import validate_office_archive

MAX_PROFILE_BYTES = 16 * 1024 * 1024


def _schema(name: str) -> dict[str, Any]:
    return json.loads((resources.files("prodocux_kernel.schemas") / name).read_text(encoding="utf-8"))


def _registry() -> Registry:
    registry = Registry()
    for name in (
        "prodocux_document_structure_profile_v1.json",
        "prodocux_template_conformance_request_v1.json",
        "prodocux_template_conformance_result_v1.json",
    ):
        schema = _schema(name)
        resource = Resource.from_contents(schema)
        registry = registry.with_resource(name, resource)
        registry = registry.with_resource(schema["$id"], resource)
    return registry


def _validate(name: str, value: Mapping[str, Any]) -> None:
    errors = sorted(
        Draft202012Validator(_schema(name), registry=_registry()).iter_errors(value),
        key=lambda item: list(item.absolute_path),
    )
    if errors:
        raise ValueError("invalid conformance contract: " + "; ".join(error.message for error in errors))


def _canonical_digest(value: Any) -> str:
    encoded = json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _column_name(index: int) -> str:
    value = index + 1
    result = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _alignment_name(value: Any) -> str:
    return {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }.get(value, "unspecified")


def _column_alignment(table: Any, column: int) -> str:
    values = {
        _alignment_name(paragraph.alignment)
        for row in table.rows
        if column < len(row.cells)
        for paragraph in row.cells[column].paragraphs
    }
    values.discard("unspecified")
    if not values:
        return "unspecified"
    return next(iter(values)) if len(values) == 1 else "mixed"


def _merge_ranges(table: Any) -> list[str]:
    positions: dict[int, list[tuple[int, int]]] = defaultdict(list)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            positions[id(cell._tc)].append((row_index, column_index))
    ranges: list[str] = []
    for cells in positions.values():
        rows = [item[0] for item in cells]
        columns = [item[1] for item in cells]
        if len(cells) <= 1:
            continue
        start = f"{_column_name(min(columns))}{min(rows) + 1}"
        end = f"{_column_name(max(columns))}{max(rows) + 1}"
        ranges.append(f"{start}:{end}")
    return sorted(set(ranges))


def profile_docx_table_structure_bytes(raw: bytes) -> dict[str, Any]:
    """Profile bounded DOCX table structure without semantic interpretation."""
    if len(raw) > MAX_PROFILE_BYTES:
        raise ValueError(f"DOCX exceeds {MAX_PROFILE_BYTES} bytes")
    validate_office_archive(raw, label="DOCX", invalid_message="invalid DOCX document")
    document = Document(BytesIO(raw))
    tables: list[dict[str, Any]] = []
    for ordinal, table in enumerate(document.tables):
        column_count = max((len(row.cells) for row in table.rows), default=0)
        if column_count < 1:
            continue
        header_rows = [
            index
            for index, row in enumerate(table.rows)
            if row._tr.get_or_add_trPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is not None
        ]
        header_index = header_rows[0] if header_rows else 0
        header_cells = table.rows[header_index].cells if table.rows else []
        columns = []
        for column in range(column_count):
            item: dict[str, Any] = {
                "ordinal": column,
                "header": header_cells[column].text.strip() if column < len(header_cells) else "",
                "alignment": _column_alignment(table, column),
                "text_direction": "unspecified",
            }
            if column < len(table.columns) and table.columns[column].width is not None:
                item["width_points"] = round(table.columns[column].width.pt, 4)
            columns.append(item)
        style_name = table.style.name if table.style is not None else ""
        tables.append({
            "table_id": f"table:{ordinal + 1:04d}", "ordinal": ordinal,
            "row_count": len(table.rows), "column_count": column_count,
            "header_rows": header_rows, "merge_ranges": _merge_ranges(table),
            "columns": columns, "fixed_row_labels": [row.cells[0].text.strip() for row in table.rows[1:] if row.cells and row.cells[0].text.strip()],
            "style": {"style_name": style_name},
        })
    result = {
        "schema_version": "prodocux_document_structure_profile_v1",
        "source_sha256": hashlib.sha256(raw).hexdigest(),
        "document_kind": "docx", "tables": tables,
        "normalization_report": [], "interpretation": "none",
    }
    _validate("prodocux_document_structure_profile_v1.json", result)
    return result


def profile_docx_table_structure(path: str | Path) -> dict[str, Any]:
    return profile_docx_table_structure_bytes(Path(path).read_bytes())


def _openpyxl_alignment(value: str | None) -> str:
    return {"left": "left", "center": "center", "right": "right", "justify": "justify"}.get(value or "", "unspecified")


def profile_xlsx_table_structure_bytes(raw: bytes) -> dict[str, Any]:
    """Profile each XLSX worksheet used range as a deterministic table."""
    if len(raw) > MAX_PROFILE_BYTES:
        raise ValueError(f"XLSX exceeds {MAX_PROFILE_BYTES} bytes")
    validate_office_archive(raw, label="XLSX", invalid_message="invalid XLSX workbook")
    workbook = load_workbook(BytesIO(raw), read_only=False, data_only=False)
    tables: list[dict[str, Any]] = []
    try:
        for ordinal, sheet in enumerate(workbook.worksheets):
            column_count = max(1, sheet.max_column)
            columns = []
            for column in range(1, column_count + 1):
                letter = _column_name(column - 1)
                width = sheet.column_dimensions[letter].width
                alignment_values = {
                    _openpyxl_alignment(sheet.cell(row=row, column=column).alignment.horizontal)
                    for row in range(1, sheet.max_row + 1)
                }
                alignment_values.discard("unspecified")
                alignment = next(iter(alignment_values)) if len(alignment_values) == 1 else ("mixed" if alignment_values else "unspecified")
                columns.append({
                    "ordinal": column - 1,
                    "header": str(sheet.cell(row=1, column=column).value or "").strip(),
                    "width_points": round(float(width or 13.0) * 7.0, 4),
                    "alignment": alignment,
                    "text_direction": "unspecified",
                })
            style_values = [
                f"{cell.style_id}:{cell.fill.patternType}:{cell.border.left.style}:{cell.border.right.style}:{cell.border.top.style}:{cell.border.bottom.style}"
                for row in sheet.iter_rows()
                for cell in row
            ]
            tables.append({
                "table_id": f"sheet:{ordinal + 1:04d}", "ordinal": ordinal,
                "anchor": sheet.title, "row_count": sheet.max_row,
                "column_count": column_count, "header_rows": [0],
                "merge_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
                "columns": columns,
                "fixed_row_labels": [str(sheet.cell(row=row, column=1).value).strip() for row in range(2, sheet.max_row + 1) if sheet.cell(row=row, column=1).value not in (None, "")],
                "style": {"style_name": "worksheet", "border_digest": _canonical_digest(style_values), "shading_digest": _canonical_digest([cell.fill.fgColor.rgb for row in sheet.iter_rows() for cell in row])},
            })
    finally:
        workbook.close()
    result = {"schema_version": "prodocux_document_structure_profile_v1", "source_sha256": hashlib.sha256(raw).hexdigest(), "document_kind": "xlsx", "tables": tables, "normalization_report": [], "interpretation": "none"}
    _validate("prodocux_document_structure_profile_v1.json", result)
    return result


def profile_xlsx_table_structure(path: str | Path) -> dict[str, Any]:
    return profile_xlsx_table_structure_bytes(Path(path).read_bytes())


def _pptx_alignment(table: Any, column: int) -> str:
    values = {
        _alignment_name(paragraph.alignment)
        for row in table.rows
        for paragraph in row.cells[column].text_frame.paragraphs
    }
    values.discard("unspecified")
    return next(iter(values)) if len(values) == 1 else ("mixed" if values else "unspecified")


def profile_pptx_table_structure_bytes(raw: bytes) -> dict[str, Any]:
    """Profile PPTX table shapes with slide/shape-stable locators."""
    if len(raw) > 32 * 1024 * 1024:
        raise ValueError("PPTX exceeds 33554432 bytes")
    validate_office_archive(raw, label="PPTX", invalid_message="invalid PPTX presentation")
    presentation = Presentation(BytesIO(raw))
    tables: list[dict[str, Any]] = []
    ordinal = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if not getattr(shape, "has_table", False):
                continue
            table = shape.table
            column_count = len(table.columns)
            tables.append({
                "table_id": f"slide:{slide_index}:table:{shape_index}", "ordinal": ordinal,
                "anchor": f"slide:{slide_index}/shape:{shape_index}",
                "row_count": len(table.rows), "column_count": column_count,
                "header_rows": [0], "merge_ranges": _merge_ranges(table),
                "columns": [{
                    "ordinal": column,
                    "header": table.cell(0, column).text.strip() if table.rows else "",
                    "width_points": round(table.columns[column].width.pt, 4),
                    "alignment": _pptx_alignment(table, column),
                    "text_direction": "unspecified",
                } for column in range(column_count)],
                "fixed_row_labels": [table.cell(row, 0).text.strip() for row in range(1, len(table.rows)) if table.cell(row, 0).text.strip()],
                "style": {"style_name": "pptx-table"},
            })
            ordinal += 1
    result = {"schema_version": "prodocux_document_structure_profile_v1", "source_sha256": hashlib.sha256(raw).hexdigest(), "document_kind": "pptx", "tables": tables, "normalization_report": [], "interpretation": "none"}
    _validate("prodocux_document_structure_profile_v1.json", result)
    return result


def profile_pptx_table_structure(path: str | Path) -> dict[str, Any]:
    return profile_pptx_table_structure_bytes(Path(path).read_bytes())


def _display(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def compare_template_conformance(request: Mapping[str, Any]) -> dict[str, Any]:
    """Compare frozen protected table invariants and return stable issues."""
    _validate("prodocux_template_conformance_request_v1.json", request)
    reference, candidate, policy = request["reference"], request["candidate"], request["policy"]
    protected = set(policy["protected"])
    issues: list[dict[str, Any]] = []

    def issue(code: str, location: str, expected: Any, actual: Any) -> None:
        issues.append({"code": code, "location": location, "message": code.replace("_", " ").lower(), "expected": _display(expected), "actual": _display(actual)})

    reference_tables = reference["tables"]
    candidate_tables = candidate["tables"]
    if "table_count" in protected and len(reference_tables) != len(candidate_tables):
        issue("TABLE_MISSING" if len(candidate_tables) < len(reference_tables) else "UNEXPECTED_TABLE_ADDED", "tables", len(reference_tables), len(candidate_tables))
    matching = policy["table_matching"]
    if matching == "table_id":
        candidate_by_key = {table["table_id"]: table for table in candidate_tables}
        pairs = [(table, candidate_by_key.get(table["table_id"])) for table in reference_tables]
    else:
        pairs = [(table, candidate_tables[index] if index < len(candidate_tables) else None) for index, table in enumerate(reference_tables)]
    for expected, actual in pairs:
        location = expected["table_id"]
        if actual is None:
            if not any(item["code"] == "TABLE_MISSING" and item["location"] == location for item in issues):
                issue("TABLE_MISSING", location, expected["table_id"], None)
            continue
        if "table_order" in protected and expected["ordinal"] != actual["ordinal"]:
            issue("TABLE_ORDER_CHANGED", location, expected["ordinal"], actual["ordinal"])
        if "topology" in protected:
            allowed_rows = policy.get("maximum_appended_rows", 0) if "append_rows" in policy["allowed"] else 0
            if actual["row_count"] < expected["row_count"] or actual["row_count"] > expected["row_count"] + allowed_rows:
                issue("ROW_COUNT_MISMATCH", location, expected["row_count"], actual["row_count"])
            if expected["column_count"] != actual["column_count"]:
                issue("COLUMN_COUNT_MISMATCH", location, expected["column_count"], actual["column_count"])
        expected_headers = [item.get("header", "") for item in expected["columns"]]
        actual_headers = [item.get("header", "") for item in actual["columns"]]
        if "headers" in protected and expected_headers != actual_headers:
            issue("HEADER_CHANGED", f"{location}/headers", expected_headers, actual_headers)
        if "merge_ranges" in protected and expected["merge_ranges"] != actual["merge_ranges"]:
            issue("MERGE_TOPOLOGY_CHANGED", f"{location}/merge_ranges", expected["merge_ranges"], actual["merge_ranges"])
        if "column_widths" in protected:
            tolerance = policy.get("column_width_tolerance_points", 0)
            for index, expected_column in enumerate(expected["columns"]):
                actual_column = actual["columns"][index] if index < len(actual["columns"]) else {}
                if "width_points" in expected_column and ("width_points" not in actual_column or abs(expected_column["width_points"] - actual_column["width_points"]) > tolerance):
                    issue("COLUMN_WIDTH_OUT_OF_TOLERANCE", f"{location}/columns/{index}/width_points", expected_column.get("width_points"), actual_column.get("width_points"))
        for field, code in (("alignment", "ALIGNMENT_CHANGED"), ("text_direction", "TEXT_DIRECTION_CHANGED")):
            if field in protected:
                expected_values = [item.get(field, "unspecified") for item in expected["columns"]]
                actual_values = [item.get(field, "unspecified") for item in actual["columns"]]
                if expected_values != actual_values:
                    issue(code, f"{location}/columns/{field}", expected_values, actual_values)
        if "style" in protected and expected["style"] != actual["style"]:
            issue("STYLE_CHANGED", f"{location}/style", expected["style"], actual["style"])
        if "fixed_row_labels" in protected:
            for label in expected.get("fixed_row_labels", []):
                if label not in actual.get("fixed_row_labels", []):
                    issue("REQUIRED_ROW_LABEL_MISSING", f"{location}/fixed_row_labels", label, None)
    result = {
        "schema_version": "prodocux_template_conformance_result_v1",
        "request_id": request["request_id"], "reference_sha256": reference["source_sha256"],
        "candidate_sha256": candidate["source_sha256"], "policy_digest": _canonical_digest(policy),
        "conforms": not issues, "issues": issues, "interpretation": "none",
    }
    _validate("prodocux_template_conformance_result_v1.json", result)
    return result
