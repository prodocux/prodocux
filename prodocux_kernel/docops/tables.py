"""docops/tables: deterministic operations on table sizing and inserting images into cells.

Handles structural issues such as images placed in tables, appendix figures,
and table overflow; generalized into the kernel.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph

PathLike = Union[str, Path]

TABLE_POLICIES = (
    "fit_to_page",
    "reuse_template_table_shape_when_possible",
)


def normalize_table_label(text: str) -> str:
    """Normalization for row-label matching: strips whitespace and full-width spaces."""
    cleaned = text.replace("\u3000", " ").replace("\n", " ")
    return re.sub(r"\s+", "", cleaned.strip())


def should_fit_tables(table_policy: str, *, fit_tables: bool = True) -> bool:
    """Decide whether to run fit_tables_to_page based on table_policy."""
    if not fit_tables:
        return False
    if table_policy == "reuse_template_table_shape_when_possible":
        return False
    return True


def fit_tables_to_page(doc: Document, width_pct: int = 5000) -> int:
    """Set every table's width as a percentage of page width (5000 = 100%) and remove fixed cell widths.

    Returns the number of tables processed. Prevents tables pasted from a
    wide-format source from overflowing the page.
    """
    count = 0
    for table in doc.tables:
        table.autofit = True
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), str(width_pct))
        for cell in table._tbl.iter_tcs():
            tc_pr = cell.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is not None:
                tc_pr.remove(tc_w)
        count += 1
    return count


def ensure_table_rows(table: Table, count: int) -> None:
    """Ensure the table has exactly `count` rows (adds rows if short, removes trailing rows if too many)."""
    while len(table.rows) < count:
        table.add_row()
    while len(table.rows) > count:
        row = table.rows[-1]
        row._tr.getparent().remove(row._tr)


def set_cell_text(cell, text: str) -> None:
    """Set cell text (preserves cell/paragraph styles)."""
    value = str(text)
    if cell.paragraphs:
        cell.paragraphs[0].text = value
    else:
        cell.add_paragraph(value)


def fill_table_data(
    table: Table,
    rows: List[List[str]],
    *,
    start_row: int = 0,
) -> int:
    """Fill in 2D data row by row starting at start_row. Returns the number of cells filled."""
    filled = 0
    for ri, row_data in enumerate(rows):
        tr_i = start_row + ri
        if tr_i >= len(table.rows):
            break
        doc_row = table.rows[tr_i]
        for ci, value in enumerate(row_data):
            if ci >= len(doc_row.cells):
                break
            set_cell_text(doc_row.cells[ci], value)
            filled += 1
    return filled


def fill_table_by_row_labels(
    table: Table,
    values: Dict[str, str],
    *,
    label_col: int = 0,
    value_col: int = 1,
    start_row: int = 0,
) -> int:
    """Keep the template's left-column labels, writing only the right column by matching labels. Returns the number of cells filled."""
    if not values:
        return 0
    index = {normalize_table_label(k): str(v) for k, v in values.items()}
    filled = 0
    for ri in range(start_row, len(table.rows)):
        row = table.rows[ri]
        if label_col >= len(row.cells) or value_col >= len(row.cells):
            continue
        label = normalize_table_label(row.cells[label_col].text)
        if not label:
            continue
        value = index.get(label)
        if value is None:
            for key, candidate in index.items():
                if key in label or label in key:
                    value = candidate
                    break
        if value is None:
            continue
        set_cell_text(row.cells[value_col], value)
        filled += 1
    return filled


def clear_cell(cell) -> None:
    """Clear a cell's content (preserving tcPr), leaving a single empty paragraph."""
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    cell.add_paragraph()


def fill_cell_with_image(
    cell,
    image_path: PathLike,
    caption: Optional[str] = None,
    width_inches: float = 5.7,
    caption_style: Optional[str] = None,
    alignment: int = 1,
) -> None:
    """Clear the cell, then insert an optional caption and one image."""
    clear_cell(cell)
    if caption:
        caption_paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        if caption_style:
            try:
                caption_paragraph.style = caption_style
            except KeyError:
                pass
        caption_paragraph.add_run(caption)
        image_paragraph = cell.add_paragraph()
    else:
        image_paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    image_paragraph.alignment = alignment
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))


def _table_from_child(child, doc: Document) -> Table:
    return Table(child, doc)


def first_table_in_children(children: List, start: int, end: int, doc: Document) -> Optional[Table]:
    """Find the first table within the [start, end) range of body children."""
    for child in children[start:end]:
        if child.tag == qn("w:tbl"):
            return _table_from_child(child, doc)
    return None


def insert_table_after(paragraph: Paragraph, rows: int, cols: int) -> Table:
    """Insert a new table after the paragraph (used when the template section originally has no table)."""
    doc = paragraph._parent
    table = doc.add_table(rows=rows, cols=cols)
    tbl_el = table._tbl
    tbl_el.getparent().remove(tbl_el)
    paragraph._p.addnext(tbl_el)
    return Table(tbl_el, doc)


def next_table_after_heading(
    doc: Document,
    heading_text: str,
    stop_styles: Optional[set] = None,
) -> Optional[Table]:
    """Find the first table after a heading and before the next heading. Returns None if not found."""
    stop = stop_styles or set()
    body = list(doc.element.body.iterchildren())
    start = None
    for idx, child in enumerate(body):
        if child.tag == qn("w:p") and Paragraph(child, doc).text.strip() == heading_text:
            start = idx
            break
    if start is None:
        return None
    for child in body[start + 1:]:
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            if p.style and p.style.name in stop and p.text.strip():
                return None
        if child.tag == qn("w:tbl"):
            return _table_from_child(child, doc)
    return None


def fill_image_in_table_row_by_label(
    doc: Document,
    row_label: str,
    image_path: PathLike,
    *,
    cell_index: int = 1,
    caption: Optional[str] = None,
    width_inches: float = 3.6,
) -> bool:
    """Find a table row whose first cell contains ``row_label`` and fill ``cell_index`` with an image."""
    needle = normalize_table_label(row_label)
    for table in doc.tables:
        for row in table.rows:
            if not row.cells:
                continue
            label = normalize_table_label(row.cells[0].text)
            if needle not in label and label not in needle:
                continue
            if cell_index >= len(row.cells):
                return False
            fill_cell_with_image(
                row.cells[cell_index],
                image_path,
                caption=caption,
                width_inches=width_inches,
            )
            return True
    return False


def fill_appendix_images(
    doc: Document,
    heading_text: str,
    items: List[Tuple[str, PathLike]],
    heading_styles: Optional[set] = None,
    width_inches: float = 5.7,
) -> bool:
    """Fill (caption, image) pairs row by row into the table after a heading. Returns whether it succeeded."""
    table = next_table_after_heading(doc, heading_text, heading_styles)
    if table is None or not items:
        return False
    ensure_table_rows(table, len(items))
    for row, (caption, image_path) in zip(table.rows, items):
        fill_cell_with_image(row.cells[0], image_path, caption=caption, width_inches=width_inches)
    return True
