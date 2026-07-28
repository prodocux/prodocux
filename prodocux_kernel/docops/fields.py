"""docops/fields: deterministic operations on Word fields.

TOC field, page-number footer (PAGE/NUMPAGES), automatic paragraph numbering.
Generalized and parameterized into the kernel.
"""
from __future__ import annotations

from typing import Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .pagination import clear_paragraph
from .sections import insert_paragraph_after


def add_field(paragraph: Paragraph, instruction: str) -> None:
    """Append a Word field (begin/instrText/separate/end) to the end of a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(end)


def set_footer_page_numbers(
    doc: Document,
    prefix: str = "Page ",
    middle: str = " of ",
    suffix: str = "",
    include_total: bool = True,
    alignment: int = 1,
) -> None:
    """Set every section's footer to a "prefix PAGE middle NUMPAGES suffix" page-number field.

    Defaults to English "Page X of Y"; for Chinese, pass prefix="第 ",
    middle=" 頁 / 共 ", suffix=" 頁".
    alignment: 0 left, 1 center, 2 right.
    """
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            footer.is_linked_to_previous = False
            for extra in list(footer.paragraphs)[1:]:
                extra._element.getparent().remove(extra._element)
            paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            clear_paragraph(paragraph)
            paragraph.alignment = alignment
            paragraph.add_run(prefix)
            add_field(paragraph, "PAGE")
            if include_total:
                paragraph.add_run(middle)
                add_field(paragraph, "NUMPAGES")
            if suffix:
                paragraph.add_run(suffix)


def insert_toc_field(
    doc: Document,
    after_heading: str,
    until_heading: Optional[str] = None,
    switches: str = 'TOC \\o "1-3" \\h \\z \\u',
) -> bool:
    """Insert a TOC field after the `after_heading` paragraph, clearing old content up to `until_heading`.

    Returns whether the insertion point was found and the field inserted.
    Deterministic and idempotent (clears the old section first).
    """
    paragraphs = list(doc.paragraphs)
    start_idx = next((i for i, p in enumerate(paragraphs) if p.text.strip() == after_heading), None)
    if start_idx is None:
        return False
    end_idx = len(paragraphs)
    if until_heading is not None:
        found_end = next(
            (i for i, p in enumerate(paragraphs) if i > start_idx and p.text.strip() == until_heading),
            None,
        )
        if found_end is not None:
            end_idx = found_end

    heading = paragraphs[start_idx]
    # Clear existing paragraphs between heading and end (keep those containing sectPr, but only clear their content)
    for paragraph in reversed(paragraphs[start_idx + 1: end_idx]):
        if paragraph._element.find(".//" + qn("w:sectPr")) is not None:
            clear_paragraph(paragraph)
        else:
            paragraph._element.getparent().remove(paragraph._element)

    toc_paragraph = insert_paragraph_after(heading, "", "Normal")
    clear_paragraph(toc_paragraph)
    add_field(toc_paragraph, switches)
    return True


def has_toc_field(doc: Document) -> bool:
    body = doc.element
    if body.xpath('.//w:fldSimple[contains(@w:instr, "TOC")]'):
        return True
    for it in body.xpath(".//w:instrText"):
        if it.text and "TOC" in it.text.upper():
            return True
    return False


def set_paragraph_numbering(paragraph: Paragraph, num_id: int) -> None:
    """Apply automatic numbering with the given numId to a paragraph (pair with parts.ensure_decimal_numbering)."""
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
