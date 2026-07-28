"""C1: Document structure invariant checker (L0 gate).

Aligned with CONTRACT.md §5.1. All checks are deterministic; no LLM calls.

Invariants fall into two categories:
- Intrinsic (decidable from a single file): valid_docx, toc_is_field,
  no_blank_pages
- Comparative (requires a reference): no_orphan_section_break,
  image_removal_no_residue
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple

from docx import Document
from docx.oxml.ns import qn


# Common wordings for TOC headings
_TOC_HEADINGS = {
    "目錄", "目录", "table of contents", "contents", "目 錄",
}
# Suspected hardcoded TOC line: dot leader + page number, or trailing page number
_HARDCODED_TOC_LINE = re.compile(r"(\.{3,}\s*\d+\s*$)|(\t\s*\d+\s*$)|(\s\d{1,4}\s*$)")

# Literal error strings Word inserts when a reference is broken
_REF_ERROR_STRINGS = (
    "Error! Reference source not found",
    "錯誤! 找不到參照來源",
    "找不到參照來源",
    "找不到參考來源",
)
# Cross-reference field instructions
_REF_FIELD = re.compile(r"\s*(?:PAGEREF|NOTEREF|REF)\s+(\S+)", re.IGNORECASE)


@dataclass
class Invariant:
    id: str
    passed: Optional[bool]  # None => skipped
    status: str             # checked | skipped | error
    code: str = ""          # language-neutral message code (localized by the presentation layer)
    params: Dict[str, Any] = field(default_factory=dict)


# --------------------------------------------------------------------------
# Low-level docx parsing helpers
# --------------------------------------------------------------------------
def _open(path: str) -> Document:
    return Document(path)


def _count_page_breaks(doc: Document) -> int:
    return len(doc.element.xpath('.//w:br[@w:type="page"]'))


def _count_sections(doc: Document) -> int:
    return len(doc.element.xpath('.//w:sectPr'))


def _count_images(doc: Document) -> int:
    # covers both modern (w:drawing) and legacy (w:pict) images
    return len(doc.element.xpath('.//w:drawing')) + len(doc.element.xpath('.//w:pict'))


def _page_content_counts(doc: Document) -> List[int]:
    """Split into "pages" by explicit page breaks / sectPr, returning the content-unit count per page (heuristic)."""
    pages: List[int] = [0]
    body = doc.element.body
    for el in body.iterchildren():
        tag = el.tag
        if tag == qn("w:p"):
            # a paragraph may contain a page break
            breaks = el.xpath('.//w:br[@w:type="page"]')
            text = "".join(el.xpath(".//w:t/text()")).strip()
            has_img = bool(el.xpath(".//w:drawing")) or bool(el.xpath(".//w:pict"))
            for _ in breaks:
                pages.append(0)
            if text or has_img:
                pages[-1] += 1
            # paragraph-level section break (next page)
            if el.xpath("./w:pPr/w:sectPr"):
                pages.append(0)
        elif tag == qn("w:tbl"):
            pages[-1] += 1
        elif tag == qn("w:sectPr"):
            # final sectPr at the end of the body
            pages.append(0)
    return pages


def _count_blank_pages(doc: Document) -> int:
    pages = _page_content_counts(doc)
    # drop the single false empty page caused by the trailing sectPr
    if len(pages) > 1 and pages[-1] == 0:
        pages = pages[:-1]
    return sum(1 for c in pages if c == 0)


def _has_toc_field(doc: Document) -> bool:
    body = doc.element
    if body.xpath('.//w:fldSimple[contains(@w:instr, "TOC")]'):
        return True
    for it in body.xpath(".//w:instrText"):
        if it.text and "TOC" in it.text.upper():
            return True
    if body.xpath('.//w:docPartGallery[@w:val="Table of Contents"]'):
        return True
    return False


def _full_body_text(doc: Document) -> str:
    return "".join(doc.element.xpath(".//w:t/text()"))


def _bookmark_names(doc: Document) -> set:
    return set(doc.element.xpath(".//w:bookmarkStart/@w:name"))


def _ref_targets(doc: Document) -> List[str]:
    targets: List[str] = []
    for it in doc.element.xpath(".//w:instrText"):
        m = _REF_FIELD.match(it.text or "")
        if m:
            targets.append(m.group(1))
    for instr in doc.element.xpath(".//w:fldSimple/@w:instr"):
        m = _REF_FIELD.match(instr or "")
        if m:
            targets.append(m.group(1))
    return targets


def _looks_hardcoded_toc(doc: Document) -> bool:
    paras = [p.text or "" for p in doc.paragraphs]
    toc_idx = None
    for i, t in enumerate(paras):
        if t.strip().lower() in _TOC_HEADINGS:
            toc_idx = i
            break
    if toc_idx is None:
        return False
    # check whether the paragraphs following the heading look like TOC lines
    for t in paras[toc_idx + 1: toc_idx + 16]:
        if t.strip() and _HARDCODED_TOC_LINE.search(t):
            return True
    return False


# --------------------------------------------------------------------------
# Invariants
# --------------------------------------------------------------------------
def check_valid_docx(path: str) -> Invariant:
    try:
        _open(path)
        return Invariant("valid_docx", True, "checked", "valid_docx.ok")
    except Exception as e:  # noqa: BLE001
        return Invariant("valid_docx", False, "checked",
                         "valid_docx.cannot_open", {"error": str(e)})


def check_toc_is_field(doc: Document) -> Invariant:
    has_field = _has_toc_field(doc)
    if has_field:
        return Invariant("toc_is_field", True, "checked", "toc_is_field.field")
    if _looks_hardcoded_toc(doc):
        return Invariant("toc_is_field", False, "checked", "toc_is_field.hardcoded")
    return Invariant("toc_is_field", True, "checked", "toc_is_field.none")


def check_no_blank_pages(doc: Document) -> Invariant:
    n = _count_blank_pages(doc)
    if n == 0:
        return Invariant("no_blank_pages", True, "checked", "no_blank_pages.ok")
    return Invariant("no_blank_pages", False, "checked",
                     "no_blank_pages.found", {"n": n})


def check_cross_references(doc: Document) -> Invariant:
    # 1) literal error strings (inserted by Word when a reference is broken)
    text = _full_body_text(doc)
    for err in _REF_ERROR_STRINGS:
        if err in text:
            return Invariant("cross_references_valid", False, "checked",
                             "cross_references_valid.error_string", {"text": err})
    # 2) REF/PAGEREF pointing to a nonexistent bookmark
    bookmarks = _bookmark_names(doc)
    broken = [t for t in _ref_targets(doc) if t not in bookmarks]
    if broken:
        return Invariant("cross_references_valid", False, "checked",
                         "cross_references_valid.broken",
                         {"count": len(broken), "sample": broken[:5]})
    return Invariant("cross_references_valid", True, "checked",
                     "cross_references_valid.ok")


def check_no_orphan_section_break(doc: Document, ref: Optional[Document]) -> Invariant:
    if ref is None:
        return Invariant("no_orphan_section_break", None, "skipped",
                         "no_orphan_section_break.skipped_no_ref")
    a, b = _count_sections(doc), _count_sections(ref)
    if a == b:
        return Invariant("no_orphan_section_break", True, "checked",
                         "no_orphan_section_break.ok", {"n": a})
    return Invariant("no_orphan_section_break", False, "checked",
                     "no_orphan_section_break.changed", {"ref": b, "target": a})


def _heading_break_map(
    doc: Document,
    heading_styles: Optional[Sequence[str]] = None,
) -> Dict[Tuple[str, str], bool]:
    """{(style_name, normalized_heading): page_break_before}, only paragraphs using a heading style."""
    from .sections import normalize_heading

    styles: Optional[Set[str]] = set(heading_styles) if heading_styles else None
    result: Dict[Tuple[str, str], bool] = {}
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text or not paragraph.style:
            continue
        name = paragraph.style.name
        if styles is not None and name not in styles:
            continue
        key = (name, normalize_heading(text))
        result[key] = bool(paragraph.paragraph_format.page_break_before)
    return result


def check_heading_pagination_matches_policy(
    doc: Document,
    *,
    heading_styles: Optional[Sequence[str]] = None,
    mapped_headings: Optional[Sequence[str]] = None,
    pagination_mode: str = "mapped_sections_only",
    first_no_break: bool = True,
    break_before_styles: Optional[Sequence[str]] = None,
    chapter_break_styles: Optional[Sequence[str]] = None,
) -> Invariant:
    """Validate heading pagination per template_rules.pagination, without relying on the template's static page_break (usually all false)."""
    from . import pagination as pag

    if not heading_styles or mapped_headings is None:
        return Invariant(
            "heading_pagination_matches_policy",
            None,
            "skipped",
            "heading_pagination_matches_policy.skipped_no_config",
        )

    expected = pag.expected_heading_page_breaks(
        doc,
        heading_styles,
        mode=pagination_mode,
        first_no_break=first_no_break,
        mapped_headings=mapped_headings,
        break_before_styles=break_before_styles,
        chapter_break_styles=chapter_break_styles,
    )
    actual = _heading_break_map(doc, heading_styles)
    if chapter_break_styles:
        actual.update(_heading_break_map(doc, chapter_break_styles))

    mismatches = [
        {
            "style": key[0],
            "heading": key[1],
            "expected": expected[key],
            "actual": actual.get(key),
        }
        for key in expected
        if actual.get(key) != expected[key]
    ]

    if pagination_mode == "none":
        unexpected = [m for m in mismatches if m["actual"]]
        if unexpected:
            return Invariant(
                "heading_pagination_matches_policy",
                False,
                "checked",
                "heading_pagination_matches_policy.unexpected_breaks",
                {"count": len(unexpected), "sample": unexpected[:5], "mode": pagination_mode},
            )
        return Invariant(
            "heading_pagination_matches_policy",
            True,
            "checked",
            "heading_pagination_matches_policy.none_ok",
            {"checked": len(expected)},
        )

    if not mismatches:
        return Invariant(
            "heading_pagination_matches_policy",
            True,
            "checked",
            "heading_pagination_matches_policy.ok",
            {"checked": len(expected), "mode": pagination_mode},
        )
    return Invariant(
        "heading_pagination_matches_policy",
        False,
        "checked",
        "heading_pagination_matches_policy.mismatch",
        {"count": len(mismatches), "sample": mismatches[:5], "mode": pagination_mode},
    )


def check_pymupdf_available(require: bool = False) -> Invariant:
    """Verify whether PyMuPDF/fitz is available (needed for CPNP/safety rendering)."""
    from ..intake.pdf_images import pymupdf_available

    ok, err = pymupdf_available()
    if ok:
        return Invariant("pymupdf_available", True, "checked", "pymupdf_available.ok")
    if require:
        return Invariant(
            "pymupdf_available",
            False,
            "checked",
            "pymupdf_available.missing",
            {"error": err or "unknown"},
        )
    return Invariant(
        "pymupdf_available",
        None,
        "skipped",
        "pymupdf_available.skipped_optional",
        {"error": err or "unknown"},
    )


def check_evidence_render_complete(evidence_index: Optional[Dict[str, Any]]) -> Invariant:
    """If there is a render item with an empty path, treat evidence extraction as incomplete."""
    if not evidence_index:
        return Invariant(
            "evidence_render_complete",
            None,
            "skipped",
            "evidence_render_complete.skipped_no_index",
        )
    failed = [
        img for img in evidence_index.get("images", [])
        if img.get("method") == "render" and not img.get("path")
    ]
    if not failed:
        return Invariant(
            "evidence_render_complete",
            True,
            "checked",
            "evidence_render_complete.ok",
            {"render_count": sum(1 for i in evidence_index.get("images", []) if i.get("method") == "render")},
        )
    return Invariant(
        "evidence_render_complete",
        False,
        "checked",
        "evidence_render_complete.failed",
        {
            "count": len(failed),
            "sample": [{"id": i.get("id"), "page": i.get("page"), "error": i.get("error")} for i in failed[:5]],
        },
    )


def check_image_removal_no_residue(doc: Document, ref: Optional[Document]) -> Invariant:
    if ref is None:
        return Invariant("image_removal_no_residue", None, "skipped",
                         "image_removal_no_residue.skipped_no_ref")
    img_a, img_b = _count_images(doc), _count_images(ref)
    blank_a, blank_b = _count_blank_pages(doc), _count_blank_pages(ref)
    if img_a < img_b and blank_a > blank_b:
        return Invariant("image_removal_no_residue", False, "checked",
                         "image_removal_no_residue.residue",
                         {"img_b": img_b, "img_a": img_a,
                          "blank_b": blank_b, "blank_a": blank_a})
    return Invariant("image_removal_no_residue", True, "checked",
                     "image_removal_no_residue.ok")


def validate_structure(
    document_path: str,
    reference_path: Optional[str] = None,
    heading_styles: Optional[Sequence[str]] = None,
    pagination_policy: Optional[Any] = None,
    mapped_headings: Optional[Sequence[str]] = None,
    evidence_index: Optional[Dict[str, Any]] = None,
    require_pymupdf: bool = False,
) -> List[Invariant]:
    """Run all L0 invariants and return the list of results."""
    results: List[Invariant] = []

    valid = check_valid_docx(document_path)
    results.append(valid)
    if not valid.passed:
        # the file can't be opened, so nothing else can be checked
        for inv_id in (
            "toc_is_field",
            "no_blank_pages",
            "cross_references_valid",
            "no_orphan_section_break",
            "image_removal_no_residue",
            "heading_pagination_matches_policy",
            "pymupdf_available",
            "evidence_render_complete",
        ):
            results.append(Invariant(inv_id, None, "skipped", "skipped.valid_docx_failed"))
        return results

    doc = _open(document_path)
    ref: Optional[Document] = None
    if reference_path:
        try:
            ref = _open(reference_path)
        except Exception as e:  # noqa: BLE001
            ref = None
            results.append(Invariant("reference_load", False, "error",
                                     "reference_load.error", {"error": str(e)}))

    results.append(check_toc_is_field(doc))
    results.append(check_no_blank_pages(doc))
    results.append(check_cross_references(doc))
    results.append(check_no_orphan_section_break(doc, ref))
    results.append(check_image_removal_no_residue(doc, ref))
    if pagination_policy is not None and mapped_headings is not None:
        results.append(
            check_heading_pagination_matches_policy(
                doc,
                heading_styles=heading_styles,
                mapped_headings=mapped_headings,
                pagination_mode=getattr(pagination_policy, "mode", "mapped_sections_only"),
                first_no_break=getattr(pagination_policy, "first_no_break", True),
                break_before_styles=getattr(pagination_policy, "break_before_styles", None),
                chapter_break_styles=getattr(pagination_policy, "chapter_break_styles", None),
            )
        )
    else:
        results.append(
            Invariant(
                "heading_pagination_matches_policy",
                None,
                "skipped",
                "heading_pagination_matches_policy.skipped_no_config",
            )
        )
    results.append(check_pymupdf_available(require=require_pymupdf))
    results.append(check_evidence_render_complete(evidence_index))
    return results


def overall_passed(invariants: List[Invariant]) -> bool:
    """Overall verdict: pass if every invariant with status=checked passes."""
    checked = [i for i in invariants if i.status == "checked"]
    return all(i.passed for i in checked) if checked else True
