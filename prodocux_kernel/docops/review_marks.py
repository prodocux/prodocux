"""docops/review_marks: highlight content that needs human confirmation (yellow/red).

Generalized into the kernel. Deterministic, no LLM calls.
"""
from __future__ import annotations

from typing import Iterable, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor


def _mark_runs(text: str, runs: Iterable, review_terms, critical_terms,
               critical_rgb: Tuple[int, int, int]) -> int:
    if not text.strip():
        return 0
    needs = any(term in text for term in review_terms)
    critical = any(term in text for term in critical_terms)
    if not needs and not critical:
        return 0
    for run in runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if critical:
            run.font.color.rgb = RGBColor(*critical_rgb)
    return 1


def highlight_review_terms(
    doc: Document,
    review_terms: List[str],
    critical_terms: Optional[List[str]] = None,
    critical_rgb: Tuple[int, int, int] = (192, 0, 0),
) -> int:
    """Highlight in yellow any paragraph/cell containing review_terms; those containing critical_terms also get red text.

    Returns the number of paragraphs marked. Used to flag "needs human
    confirmation" spots before delivery.
    """
    critical_terms = critical_terms or []
    marked = 0
    for paragraph in doc.paragraphs:
        marked += _mark_runs(paragraph.text, paragraph.runs, review_terms, critical_terms, critical_rgb)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    marked += _mark_runs(paragraph.text, paragraph.runs, review_terms,
                                         critical_terms, critical_rgb)
    return marked


def clear_highlights(doc: Document) -> int:
    """Clear highlight marks throughout the document (paragraphs and tables). Returns the number of runs cleared."""
    count = 0
    def clear(paragraph):
        nonlocal count
        for run in paragraph.runs:
            if run.font.highlight_color is not None:
                run.font.highlight_color = None
                count += 1
    for paragraph in doc.paragraphs:
        clear(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    clear(paragraph)
    return count
