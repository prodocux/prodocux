"""Flagship doc-ops library — structural-fix regression tests.

Fully deterministic. Each test locks down a structural issue that has
recurred in the past.
"""
import struct
import zlib

import pytest
from docx import Document
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
from docx.oxml.ns import qn

from prodocux_kernel.docops import (
    fields,
    front_matter,
    pagination,
    parts,
    review_marks,
    sections,
    tables,
)
from prodocux_kernel.docops.front_matter import FrontMatterAnchor


def _png_bytes(width: int = 1, height: int = 1) -> bytes:
    """Programmatically generate a valid PNG (avoids mistakes in hardcoded base64)."""
    def chunk(typ: bytes, data: bytes) -> bytes:
        body = typ + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)  # 8-bit RGB
    raw = (b"\x00" + b"\xff\xff\xff" * width) * height
    idat = zlib.compress(raw)
    return sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


@pytest.fixture
def png(tmp_path):
    p = tmp_path / "img.png"
    p.write_bytes(_png_bytes())
    return str(p)


# ---------------- parts (zip layer) ----------------
def test_enable_update_fields_on_open(tmp_path):
    doc = Document()
    doc.add_paragraph("x")
    path = tmp_path / "a.docx"
    doc.save(str(path))
    parts.enable_update_fields_on_open(path)
    settings = parts.read_part(path, "word/settings.xml").decode("utf-8")
    assert "w:updateFields" in settings
    # idempotent: calling again does not insert a duplicate
    parts.enable_update_fields_on_open(path)
    assert parts.read_part(path, "word/settings.xml").decode("utf-8").count("w:updateFields") == 1


def test_ensure_decimal_numbering(tmp_path):
    doc = Document()
    doc.add_paragraph("item", style="List Number")  # triggers numbering.xml
    path = tmp_path / "n.docx"
    doc.save(str(path))
    if not parts.has_part(path, "word/numbering.xml"):
        pytest.skip("This environment's default template has no numbering.xml")
    parts.ensure_decimal_numbering(path, 900)
    text = parts.read_part(path, "word/numbering.xml").decode("utf-8")
    assert 'w:numId="900"' in text


# ---------------- fields ----------------
def test_set_footer_page_numbers(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    fields.set_footer_page_numbers(doc, prefix="第 ", middle=" 頁 / 共 ", suffix=" 頁")
    xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in xml and "NUMPAGES" in xml


def test_insert_toc_field(tmp_path):
    doc = Document()
    doc.add_paragraph("目錄")
    doc.add_paragraph("舊的假目錄行 1")
    doc.add_paragraph("產品敘述", style="Heading 1")
    ok = fields.insert_toc_field(doc, after_heading="目錄", until_heading="產品敘述")
    assert ok is True
    assert fields.has_toc_field(doc) is True


def test_set_paragraph_numbering():
    doc = Document()
    p = doc.add_paragraph("標題", style="Heading 1")
    fields.set_paragraph_numbering(p, 900)
    num_prs = p._p.xpath(".//w:numPr")
    assert num_prs
    assert num_prs[0].find(qn("w:numId")).get(qn("w:val")) == "900"


# ---------------- pagination (blank pages / page breaks) ----------------
def test_remove_blank_page_breaks_before_headings():
    doc = Document()
    blank = doc.add_paragraph()
    blank.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("章節", style="Heading 1")
    removed = pagination.remove_blank_page_breaks_before_headings(doc, {"Heading 1"})
    assert removed == 1
    assert all(p.text.strip() or "章節" in p.text for p in doc.paragraphs)


def test_set_heading_page_breaks_all_heading_styles():
    doc = Document()
    for t in ("一", "二", "三"):
        doc.add_paragraph(t, style="Heading 1")
    count = pagination.set_heading_page_breaks(
        doc, {"Heading 1"}, first_no_break=True, mode="all_heading_styles",
    )
    assert count == 2
    headings = [p for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
    assert not headings[0].paragraph_format.page_break_before
    assert headings[1].paragraph_format.page_break_before is True


def test_set_heading_page_breaks_mapped_sections_only_skips_foreword():
    """An unmapped front-matter heading must not consume first_no_break; the first mapped section should not page-break."""
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    for name in ("PIF標題一", "PIF標題二"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("產品敘述", style="PIF標題一")
    doc.add_paragraph("前言內文")
    doc.add_paragraph("產品基本資料", style="PIF標題二")
    doc.add_paragraph("完成產品登錄之證明文件", style="PIF標題二")
    count = pagination.set_heading_page_breaks(
        doc,
        {"PIF標題一", "PIF標題二"},
        first_no_break=True,
        mode="mapped_sections_only",
        mapped_headings=["產品基本資料", "完成產品登錄之證明文件"],
    )
    assert count == 1
    by_text = {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}
    assert not by_text["產品敘述"].paragraph_format.page_break_before
    assert not by_text["產品基本資料"].paragraph_format.page_break_before
    assert by_text["完成產品登錄之證明文件"].paragraph_format.page_break_before is True


def test_apply_pagination_policy_tw_pif_like():
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    for name in ("PIF標題一", "PIF標題二"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("產品敘述", style="PIF標題一")
    doc.add_paragraph("產品基本資料", style="PIF標題二")
    doc.add_paragraph("完成產品登錄之證明文件", style="PIF標題二")
    doc.add_paragraph("品質資料", style="PIF標題一")

    count = pagination.apply_pagination_policy(
        doc,
        {"PIF標題一", "PIF標題二"},
        mode="mapped_sections_only",
        first_no_break=True,
        mapped_headings=["產品基本資料", "完成產品登錄之證明文件"],
        chapter_break_styles=["PIF標題一"],
    )
    assert count == 2
    by_text = {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}
    assert not by_text["產品敘述"].paragraph_format.page_break_before
    assert not by_text["產品基本資料"].paragraph_format.page_break_before
    assert by_text["完成產品登錄之證明文件"].paragraph_format.page_break_before is True
    assert by_text["品質資料"].paragraph_format.page_break_before is True


def test_set_heading_page_breaks_none_is_noop():
    doc = Document()
    doc.add_paragraph("A", style="Heading 1")
    doc.add_paragraph("B", style="Heading 1")
    assert pagination.set_heading_page_breaks(doc, {"Heading 1"}, mode="none") == 0
    headings = [p for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
    assert not headings[0].paragraph_format.page_break_before
    assert not headings[1].paragraph_format.page_break_before


# ---------------- sections (heading-anchored replacement) ----------------
def test_replace_sections():
    doc = Document()
    doc.add_paragraph("Section A", style="Heading 1")
    doc.add_paragraph("old a1")
    doc.add_paragraph("old a2")
    doc.add_paragraph("Section B", style="Heading 1")
    doc.add_paragraph("keep b1")

    headings = {
        sections.normalize_heading("Section A"): "A",
        sections.normalize_heading("Section B"): "B",
    }
    replaced = sections.replace_sections(
        doc, {"A": ["new a"]}, headings, {"Heading 1"}, body_style=None
    )
    assert replaced == 1
    texts = [p.text for p in doc.paragraphs]
    assert "new a" in texts
    assert "old a1" not in texts and "old a2" not in texts
    assert "keep b1" in texts  # other sections are unaffected


def test_replace_last_section_preserves_trailing_sectpr():
    """Replacing "the last section in the document" must not remove the body's trailing sectPr (or it would break sectioning)."""
    doc = Document()
    doc.add_paragraph("Only Section", style="Heading 1")
    doc.add_paragraph("old tail")
    before = len(doc.element.xpath(".//w:sectPr"))
    headings = {sections.normalize_heading("Only Section"): "S"}
    sections.replace_sections(doc, {"S": ["new tail"]}, headings, {"Heading 1"})
    after = len(doc.element.xpath(".//w:sectPr"))
    assert after == before == 1
    texts = [p.text for p in doc.paragraphs]
    assert "new tail" in texts and "old tail" not in texts


def test_insert_paragraph_after():
    doc = Document()
    anchor = doc.add_paragraph("anchor")
    sections.insert_paragraph_after(anchor, "inserted", None)
    texts = [p.text for p in doc.paragraphs]
    assert texts[texts.index("anchor") + 1] == "inserted"


def test_remove_images_in_section(png):
    doc = Document()
    doc.add_paragraph("圖區", style="Heading 1")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).paragraphs[0].add_run().add_picture(png)
    doc.add_paragraph("下一節", style="Heading 1")
    removed = sections.remove_images_in_section(doc, "圖區", {"Heading 1"})
    assert removed == 1


# ---------------- tables (images / overflow) ----------------
def test_should_fit_tables_respects_policy():
    assert tables.should_fit_tables("fit_to_page", fit_tables=True) is True
    assert tables.should_fit_tables("reuse_template_table_shape_when_possible", fit_tables=True) is False
    assert tables.should_fit_tables("fit_to_page", fit_tables=False) is False


def test_fit_tables_to_page():
    doc = Document()
    doc.add_table(rows=2, cols=3)
    count = tables.fit_tables_to_page(doc)
    assert count == 1
    tbl_w = doc.tables[0]._tbl.tblPr.find(qn("w:tblW"))
    assert tbl_w.get(qn("w:type")) == "pct"
    assert tbl_w.get(qn("w:w")) == "5000"


def test_fill_table_by_row_labels_preserves_left_column():
    doc = Document()
    t = doc.add_table(rows=3, cols=2)
    tables.set_cell_text(t.cell(0, 0), "項目")
    tables.set_cell_text(t.cell(0, 1), "內容描述")
    tables.set_cell_text(t.cell(1, 0), "產品名稱(中文/英文)")
    tables.set_cell_text(t.cell(1, 1), "OLD")
    tables.set_cell_text(t.cell(2, 0), "產品類別")
    tables.set_cell_text(t.cell(2, 1), "OLD2")
    filled = tables.fill_table_by_row_labels(
        t,
        {"產品名稱(中文/英文)": "Demo Perfume 示範香水", "產品類別": "香水（一般化粧品）"},
        start_row=1,
    )
    assert filled == 2
    assert t.cell(1, 0).text == "產品名稱(中文/英文)"
    assert "Demo Perfume" in t.cell(1, 1).text
    assert t.cell(2, 0).text == "產品類別"
    assert t.cell(2, 1).text == "香水（一般化粧品）"


def test_ensure_table_rows():
    doc = Document()
    t = doc.add_table(rows=1, cols=2)
    tables.ensure_table_rows(t, 3)
    assert len(t.rows) == 3
    tables.ensure_table_rows(t, 2)
    assert len(t.rows) == 2


def test_fill_cell_with_image(png):
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    tables.fill_cell_with_image(t.cell(0, 0), png, caption="圖一", width_inches=2.0)
    blips = t.cell(0, 0)._tc.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    )
    assert len(blips) == 1
    assert "圖一" in t.cell(0, 0).text


def test_fill_appendix_images(png):
    doc = Document()
    doc.add_paragraph("附錄圖", style="Heading 1")
    doc.add_table(rows=1, cols=1)
    ok = tables.fill_appendix_images(
        doc, "附錄圖", [("第1張", png), ("第2張", png)], heading_styles={"Heading 1"}
    )
    assert ok is True
    assert len(doc.tables[0].rows) == 2


# ---------------- front matter ----------------
def test_apply_front_matter_by_style_and_prefix():
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    doc.styles.add_style("CoverMeta", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("DEMO BRAND", style="Title")
    doc.add_paragraph("Old Product", style="Normal")
    doc.add_paragraph("文件編號：OLD-ID", style="CoverMeta")
    doc.add_paragraph("文件版本：1.0", style="CoverMeta")
    anchors = [
        FrontMatterAnchor(field="product_name_display", style="Normal", occurrence=1),
        FrontMatterAnchor(field="document_id", style="CoverMeta", prefix="文件編號："),
    ]
    written = front_matter.apply_front_matter(
        doc,
        anchors,
        {"product_name_display": "Demo Perfume 示範香水", "document_id": "PIF-DEMO-2026-001"},
        stop_styles={"PIF標題一"},
    )
    assert written == 2
    assert doc.paragraphs[1].text == "Demo Perfume 示範香水"
    assert doc.paragraphs[2].text == "文件編號：PIF-DEMO-2026-001"
    assert doc.paragraphs[3].text == "文件版本：1.0"


# ---------------- review marks ----------------
def test_highlight_and_clear_review_terms():
    doc = Document()
    doc.add_paragraph("此處待補充內容")
    doc.add_paragraph("正常內容")
    marked = review_marks.highlight_review_terms(doc, review_terms=["待補充"])
    assert marked == 1
    run = doc.paragraphs[0].runs[0]
    assert run.font.highlight_color == WD_COLOR_INDEX.YELLOW
    cleared = review_marks.clear_highlights(doc)
    assert cleared == 1
