"""C1 invariant checker tests. Generates docx fixtures programmatically."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from prodocux_kernel.docops import invariants as inv


def _add_fake_drawing(paragraph) -> None:
    """Inject a w:drawing element to simulate an embedded image (for image-counting tests)."""
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:drawing"))
    paragraph._p.append(r)


def _save(doc: Document, tmp_path: Path, name: str) -> str:
    p = tmp_path / name
    doc.save(str(p))
    return str(p)


def _by_id(results, inv_id):
    return next(r for r in results if r.id == inv_id)


# --------------------------------------------------------------------------
def test_valid_docx_pass_and_fail(tmp_path):
    doc = Document()
    doc.add_paragraph("hello")
    good = _save(doc, tmp_path, "good.docx")
    assert _by_id(inv.validate_structure(good), "valid_docx").passed is True

    bad = tmp_path / "bad.docx"
    bad.write_text("not a real docx", encoding="utf-8")
    res = inv.validate_structure(str(bad))
    assert _by_id(res, "valid_docx").passed is False
    # everything else is skipped when the file can't be opened
    assert _by_id(res, "toc_is_field").status == "skipped"


def test_toc_no_toc_passes(tmp_path):
    doc = Document()
    doc.add_paragraph("一般內文，沒有目錄")
    res = inv.validate_structure(_save(doc, tmp_path, "notoc.docx"))
    assert _by_id(res, "toc_is_field").passed is True


def test_toc_field_passes(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    r = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.text = 'TOC \\o "1-3" \\h'
    r.append(instr)
    p._p.append(r)
    res = inv.validate_structure(_save(doc, tmp_path, "tocfield.docx"))
    assert _by_id(res, "toc_is_field").passed is True


def test_hardcoded_toc_fails(tmp_path):
    doc = Document()
    doc.add_paragraph("目錄")
    doc.add_paragraph("第一章 簡介 ............... 1")
    doc.add_paragraph("第二章 方法 ............... 5")
    doc.add_paragraph("正文開始")
    res = inv.validate_structure(_save(doc, tmp_path, "hardtoc.docx"))
    assert _by_id(res, "toc_is_field").passed is False


def test_no_blank_pages_pass(tmp_path):
    doc = Document()
    doc.add_paragraph("A")
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("B")
    res = inv.validate_structure(_save(doc, tmp_path, "clean.docx"))
    assert _by_id(res, "no_blank_pages").passed is True


def test_blank_page_detected(tmp_path):
    doc = Document()
    doc.add_paragraph("A")
    # two consecutive page breaks with no content between them => produces a blank page
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)
    run.add_break(WD_BREAK.PAGE)
    doc.add_paragraph("B")
    res = inv.validate_structure(_save(doc, tmp_path, "blank.docx"))
    assert _by_id(res, "no_blank_pages").passed is False


def test_section_break_comparative(tmp_path):
    ref = Document()
    ref.add_paragraph("A")
    ref.add_section()
    ref_path = _save(ref, tmp_path, "ref_sec.docx")

    target = Document()
    target.add_paragraph("A")
    target_path = _save(target, tmp_path, "tgt_sec.docx")

    res = inv.validate_structure(target_path, ref_path)
    assert _by_id(res, "no_orphan_section_break").passed is False

    # matching section count should pass
    res2 = inv.validate_structure(ref_path, ref_path)
    assert _by_id(res2, "no_orphan_section_break").passed is True


def test_image_removal_residue(tmp_path):
    ref = Document()
    ref.add_paragraph("A")
    _add_fake_drawing(ref.add_paragraph())
    ref.add_paragraph("B")
    ref_path = _save(ref, tmp_path, "ref_img.docx")

    # target: the image was removed but the page breaks before and after remain => produces a blank page
    target = Document()
    target.add_paragraph("A")
    run = target.add_paragraph().add_run()
    run.add_break(WD_BREAK.PAGE)
    run.add_break(WD_BREAK.PAGE)
    target.add_paragraph("B")
    target_path = _save(target, tmp_path, "tgt_img.docx")

    res = inv.validate_structure(target_path, ref_path)
    assert _by_id(res, "image_removal_no_residue").passed is False


def test_cross_reference_error_string_fails(tmp_path):
    doc = Document()
    doc.add_paragraph("詳見 錯誤! 找不到參照來源。")
    res = inv.validate_structure(_save(doc, tmp_path, "xref_err.docx"))
    assert _by_id(res, "cross_references_valid").passed is False


def test_cross_reference_clean_passes(tmp_path):
    doc = Document()
    doc.add_paragraph("一般內文，沒有交叉參照")
    res = inv.validate_structure(_save(doc, tmp_path, "xref_ok.docx"))
    assert _by_id(res, "cross_references_valid").passed is True


def test_broken_ref_field_fails(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    r = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGEREF _Ref_missing \\h"
    r.append(instr)
    p._p.append(r)
    res = inv.validate_structure(_save(doc, tmp_path, "ref_broken.docx"))
    assert _by_id(res, "cross_references_valid").passed is False


def test_comparative_skipped_without_reference(tmp_path):
    doc = Document()
    doc.add_paragraph("A")
    res = inv.validate_structure(_save(doc, tmp_path, "x.docx"))
    assert _by_id(res, "no_orphan_section_break").status == "skipped"
    assert _by_id(res, "image_removal_no_residue").status == "skipped"
    assert _by_id(res, "heading_pagination_matches_policy").status == "skipped"


def test_heading_pagination_policy_detects_wrong_first_mapped_break():
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    for name in ("PIF標題一", "PIF標題二"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("產品敘述", style="PIF標題一")
    p = doc.add_paragraph("產品基本資料", style="PIF標題二")
    p.paragraph_format.page_break_before = True

    res = inv.check_heading_pagination_matches_policy(
        doc,
        heading_styles=["PIF標題一", "PIF標題二"],
        mapped_headings=["產品基本資料", "完成產品登錄之證明文件"],
        pagination_mode="mapped_sections_only",
        first_no_break=True,
        chapter_break_styles=["PIF標題一"],
    )
    assert res.passed is False

    res_ok = inv.check_heading_pagination_matches_policy(
        doc,
        heading_styles=["PIF標題一", "PIF標題二"],
        mapped_headings=["產品基本資料"],
        pagination_mode="none",
        first_no_break=True,
    )
    assert res_ok.passed is False
    assert res_ok.code == "heading_pagination_matches_policy.unexpected_breaks"
