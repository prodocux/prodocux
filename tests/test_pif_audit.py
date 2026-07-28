"""Skill #9 Taiwan PIF compliance audit tests."""
import os
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from prodocux_kernel.audit.pif_tw import audit_pif_tw, load_checklist
from prodocux_kernel.templating.mapping import MappingConfig
from skills.pif_audit import pif_audit as skill

EXAMPLES = Path(__file__).resolve().parents[1] / "examples" / "pif_tw"
CHECKLIST = EXAMPLES / "tw_pif_checklist.json"
MAPPING = EXAMPLES / "template_mapping.json"
# Optional local template — set PRODOCUX_TW_PIF_TEMPLATE; never hardcode customer filenames in-repo.
_REAL_TPL = os.environ.get("PRODOCUX_TW_PIF_TEMPLATE", "").strip()
REAL_TPL = Path(_REAL_TPL) if _REAL_TPL else Path("__unset_PRODOCUX_TW_PIF_TEMPLATE__")


def _mapping():
    return MappingConfig.from_json(MAPPING)


def _checklist():
    return load_checklist(CHECKLIST)


def _fill_line(sid: str) -> str:
    """One paragraph of substantive content per section, containing the required keywords."""
    samples = {
        "01_basic": "產品名稱：測試香水；產品類別：香水；劑型：液劑；用途：賦香；製造廠：Test Factory；輸入業者：Test Importer",
        "02_registration": "完成產品登錄之證明文件，CPNP 通報編號 DEMO-0000001",
        "03_formula": "全成分 INCI 名稱及含量：Aqua 70%, Alcohol denat. 25%, Parfum 5%, CAS 7732-18-5",
        "04_label": "產品標籤、外包裝及容器照片與標示說明",
        "05_gmp": "製造場所符合 ISO 22716 GMP 優良製造準則之聲明書",
        "06_manufacturing": "製造方法與生產流程：混合、充填、包裝",
        "07_use": "使用方法、使用部位、用量、頻率及使用族群說明",
        "08_adverse": "產品使用不良反應資料：目前無不良反應紀錄",
        "09_physical_chemical": "產品及各成分之物理及化學特性資料",
        "10_toxicology": "成分毒理資料 toxicology MoS SED NOAEL 評估",
        "11_stability": "產品安定性試驗報告 stability PAO 36 個月",
        "12_microbiology": "微生物檢測報告 microbiology 結果",
        "13_preservative": "防腐效能試驗報告 preservative challenge test",
        "14_claims": "功能評估佐證資料：香氣賦香 claim",
        "15_packaging_material": "與產品接觸之包裝材質資料 packaging material 94/62",
        "16_safety": "安全性評估結論與建議；安全資料簽署人員資格證明；safety assessment conclusion",
    }
    return samples.get(sid, f"content for {sid}")


def _pif_doc(tmp_path, *, omit: str | None = None, placeholder: str | None = None,
             review_in: str | None = None):
    config = _mapping()
    doc = Document()
    for name in (*config.heading_styles, config.body_style):
        if name not in [s.name for s in doc.styles]:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    h1, h2 = config.heading_styles
    doc.add_paragraph("目錄", style=h1)
    doc.add_paragraph("產品敘述", style=h1)
    for spec in config.sections:
        if spec.id == omit:
            continue
        doc.add_paragraph(spec.heading, style=h2)
        text = placeholder if placeholder else _fill_line(spec.id)
        if review_in == spec.id:
            text = text + " REVIEW"
        doc.add_paragraph(text, style=config.body_style)
        if spec.mode == "table":
            doc.add_table(rows=2, cols=2)
    path = tmp_path / "pif.docx"
    doc.save(str(path))
    return str(path)


def test_checklist_loads():
    cl = _checklist()
    assert len(cl["required_sections"]) == 16
    assert "化粧品產品資訊檔案管理辦法" in cl["source"]["regulation"]


def test_good_pif_passes(tmp_path):
    path = _pif_doc(tmp_path)
    result = audit_pif_tw(path, _mapping(), _checklist())
    assert result["sections_found"] == 16
    assert result["passed"] is True
    assert result["fail_count"] == 0


def test_missing_section_fails(tmp_path):
    path = _pif_doc(tmp_path, omit="11_stability")
    result = audit_pif_tw(path, _mapping(), _checklist())
    assert result["passed"] is False
    checks = {f["check"] for f in result["findings"]}
    assert "section_missing" in checks


def test_placeholder_fails(tmp_path):
    path = _pif_doc(tmp_path, placeholder="待補充")
    result = audit_pif_tw(path, _mapping(), _checklist())
    assert result["passed"] is False
    assert any(f["check"] == "section_empty_or_placeholder" for f in result["findings"])


def test_review_marker_fails(tmp_path):
    path = _pif_doc(tmp_path, review_in="02_registration")
    result = audit_pif_tw(path, _mapping(), _checklist())
    assert any(f["check"] == "review_marker_in_section" for f in result["findings"])


def test_report_bilingual(tmp_path):
    path = _pif_doc(tmp_path, omit="03_formula")
    result = audit_pif_tw(path, _mapping(), _checklist())
    en = skill.format_report(result, "en")
    zh = skill.format_report(result, "zh-TW")
    assert "Regulation:" in en and "Mandatory section missing" in en
    assert "法規依據" in zh and "缺少必要區段" in zh


@pytest.mark.skipif(not REAL_TPL.is_file(), reason="Set PRODOCUX_TW_PIF_TEMPLATE for optional template audit")
def test_optional_template_audit():
    """Optional audit against a private TW PIF template (placeholders expected to fail some checks)."""
    result = audit_pif_tw(str(REAL_TPL), _mapping(), _checklist())
    assert result["sections_found"] == 16
    # curated templates should resolve all mapped sections; content quality is separate
    assert result["sections_required"] == 16
