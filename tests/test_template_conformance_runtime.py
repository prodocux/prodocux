from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from prodocux_kernel.verification import (
    compare_template_conformance,
    profile_docx_table_structure_bytes,
    profile_pptx_table_structure_bytes,
    profile_xlsx_table_structure_bytes,
)


def _docx(*, header: str = "Shot", rows: int = 2, merge: bool = False) -> bytes:
    document = Document()
    table = document.add_table(rows=rows, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.cell(0, 0).text = header
    table.cell(0, 1).text = "Duration"
    for index in range(1, rows):
        table.cell(index, 0).text = f"shot-{index}"
        table.cell(index, 1).text = "5"
    if merge:
        table.cell(1, 0).merge(table.cell(1, 1))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _request(reference: dict, candidate: dict, *, allowed: list[str] | None = None, maximum_appended_rows: int | None = None) -> dict:
    policy = {
        "table_matching": "table_id",
        "protected": ["table_count", "table_order", "topology", "headers", "merge_ranges", "column_widths", "style", "fixed_row_labels"],
        "allowed": allowed or ["body_text"],
        "column_width_tolerance_points": 0.1,
    }
    if maximum_appended_rows is not None:
        policy["maximum_appended_rows"] = maximum_appended_rows
    return {"schema_version": "prodocux_template_conformance_request_v1", "request_id": "request:runtime:001", "reference": reference, "candidate": candidate, "policy": policy}


def test_real_docx_profiles_and_conforms() -> None:
    reference = profile_docx_table_structure_bytes(_docx())
    candidate = profile_docx_table_structure_bytes(_docx())
    assert reference["tables"][0]["columns"][0]["header"] == "Shot"
    result = compare_template_conformance(_request(reference, candidate))
    assert result["conforms"] is True
    assert result["issues"] == []


def test_header_and_merge_changes_have_stable_codes() -> None:
    reference = profile_docx_table_structure_bytes(_docx())
    candidate = profile_docx_table_structure_bytes(_docx(header="Scene", merge=True))
    result = compare_template_conformance(_request(reference, candidate))
    assert result["conforms"] is False
    assert {item["code"] for item in result["issues"]} >= {"HEADER_CHANGED", "MERGE_TOPOLOGY_CHANGED"}


def test_append_rows_respects_explicit_policy() -> None:
    reference = profile_docx_table_structure_bytes(_docx(rows=2))
    candidate = profile_docx_table_structure_bytes(_docx(rows=3))
    rejected = compare_template_conformance(_request(reference, candidate))
    assert "ROW_COUNT_MISMATCH" in {item["code"] for item in rejected["issues"]}
    accepted = compare_template_conformance(_request(reference, candidate, allowed=["body_text", "append_rows"], maximum_appended_rows=1))
    assert "ROW_COUNT_MISMATCH" not in {item["code"] for item in accepted["issues"]}


def _xlsx(*, header: str = "Shot") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Shots"
    sheet.append([header, "Duration"])
    sheet.append(["shot-1", 5])
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "notes"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_xlsx_uses_same_frozen_profile_and_comparer() -> None:
    reference = profile_xlsx_table_structure_bytes(_xlsx())
    candidate = profile_xlsx_table_structure_bytes(_xlsx(header="Scene"))
    assert reference["document_kind"] == "xlsx"
    assert reference["tables"][0]["anchor"] == "Shots"
    assert reference["tables"][0]["merge_ranges"] == ["A3:B3"]
    result = compare_template_conformance(_request(reference, candidate))
    assert "HEADER_CHANGED" in {item["code"] for item in result["issues"]}


def _pptx(*, header: str = "Shot") -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    shape.table.cell(0, 0).text = header
    shape.table.cell(0, 1).text = "Duration"
    shape.table.cell(1, 0).text = "shot-1"
    shape.table.cell(1, 1).text = "5"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_pptx_uses_same_frozen_profile_and_comparer() -> None:
    reference = profile_pptx_table_structure_bytes(_pptx())
    candidate = profile_pptx_table_structure_bytes(_pptx(header="Scene"))
    assert reference["document_kind"] == "pptx"
    assert reference["tables"][0]["anchor"] == "slide:1/shape:1"
    result = compare_template_conformance(_request(reference, candidate))
    assert "HEADER_CHANGED" in {item["code"] for item in result["issues"]}
