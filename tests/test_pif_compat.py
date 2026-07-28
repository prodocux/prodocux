"""Compatibility tests: the kernel's flagship pipeline can ingest the curated TW PIF mapping config.

- #1: loads examples/pif_tw/template_mapping.json (a public, synthetic /
  curated config).
- #2: synthesizes a template using the real styles (PIF標題一/二, PIF內文4);
  extract_profile should hit 16/16.
- #3: after render, the structure health-check gate is fully green (including
  TOC regeneration clearing stale PAGEREF).

Portable (does not depend on any local external absolute paths).
"""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from prodocux_kernel.templating import profile, render
from prodocux_kernel.templating.mapping import MappingConfig, validate_config

CONFIG_PATH = Path(__file__).resolve().parents[1] / "examples" / "pif_tw" / "template_mapping.json"


def _config():
    return MappingConfig.from_json(CONFIG_PATH)


def _synthetic_template(tmp_path, config: MappingConfig):
    """Build a synthetic template using real PIF style names (including a leftover old static TOC line)."""
    doc = Document()
    for name in (*config.heading_styles, config.body_style):
        if name not in [s.name for s in doc.styles]:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    h1, h2 = config.heading_styles[0], config.heading_styles[1]

    doc.add_paragraph("目錄", style=h1)
    doc.add_paragraph("舊目錄行 1 ........ 3", style=config.body_style)  # simulates a stale static TOC
    doc.add_paragraph("產品敘述", style=h1)                              # heading of the first content group
    for spec in config.sections:
        doc.add_paragraph(spec.heading, style=h2)
        doc.add_paragraph("(待填)", style=config.body_style)
        if spec.mode == "table":
            doc.add_table(rows=2, cols=2)
    path = tmp_path / "synthetic_pif.docx"
    doc.save(str(path))
    return str(path)


def test_pif_config_loads_and_validates():
    config = _config()
    assert config.heading_styles == ["PIF標題一", "PIF標題二"]
    assert config.body_style == "PIF內文4"
    assert len(config.sections) == 16
    assert sorted({s.mode for s in config.sections}) == ["paragraphs", "table"]
    assert validate_config(config) == []
    assert "待補充" in config.remove_review_artifacts


def test_pif_profile_extracts_all_sections(tmp_path):
    config = _config()
    tpl = _synthetic_template(tmp_path, config)
    prof = profile.extract_profile(tpl, config.section_specs(), config.heading_styles, config.body_style)
    assert prof.missing == []
    assert sum(1 for s in prof.sections if s.found) == 16
    assert all(s.match == "exact" for s in prof.sections)
    # table-mode sections should have a table detected
    formula = prof.section("03_formula")
    assert formula.has_table and (formula.table_rows, formula.table_cols) == (2, 2)


def test_pif_render_passes_gate(tmp_path):
    config = _config()
    tpl = _synthetic_template(tmp_path, config)
    out = tmp_path / "out_pif.docx"
    drafts = {
        "02_registration": ["CPNP 通報資訊草稿一", "草稿二"],
        "11_stability": ["安定性草稿"],
    }
    res = render.render_document(
        tpl, out, drafts, config,
        toc=render.TocSpec(after_heading="目錄", until_heading="產品敘述"),
        footer=render.FooterSpec(prefix="第 ", middle=" 頁 / 共 ", suffix=" 頁"),
        review_terms=["待補充"],
    )
    assert res.passed is True
    assert res.sections_written == 2
    assert res.toc_inserted is True
