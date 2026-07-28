"""#7 Number consistency audit — engine and skill tests."""
from docx import Document

from prodocux_kernel.audit import numbers as na
from skills.number_audit import number_audit as nau


def test_parse_cn_number():
    assert na.parse_cn_number("壹萬貳仟") == 12000
    assert na.parse_cn_number("拾") == 10
    assert na.parse_cn_number("貳佰零五") == 205
    assert na.parse_cn_number("壹億") == 100000000


def _doc_with_table(tmp_path, rows, name="tbl"):
    doc = Document()
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.cell(ri, ci).text = str(val)
    out = tmp_path / f"{name}.docx"
    doc.save(str(out))
    return str(out)


def test_table_total_mismatch(tmp_path):
    path = _doc_with_table(tmp_path, [
        ["項目", "金額"],
        ["A", "2000"],
        ["B", "2800"],
        ["合計", "5000"],  # should actually be 4800
    ])
    res = na.audit_numbers(path)
    checks = [f["check"] for f in res["findings"]]
    assert "table_total_mismatch" in checks
    assert res["passed"] is False


def test_table_total_correct_passes(tmp_path):
    path = _doc_with_table(tmp_path, [
        ["項目", "金額"],
        ["A", "2000"],
        ["B", "2800"],
        ["合計", "4800"],
    ])
    res = na.audit_numbers(path)
    assert all(f["check"] != "table_total_mismatch" for f in res["findings"])


def test_amount_words_mismatch(tmp_path):
    doc = Document()
    doc.add_paragraph("總價：新台幣壹萬貳仟元整（21,000）")
    out = tmp_path / "amt.docx"
    doc.save(str(out))
    res = na.audit_numbers(str(out))
    assert any(f["check"] == "amount_words_mismatch" for f in res["findings"])


def test_amount_words_match_passes(tmp_path):
    doc = Document()
    doc.add_paragraph("總價：新台幣壹萬貳仟元整（12,000）")
    out = tmp_path / "amt_ok.docx"
    doc.save(str(out))
    res = na.audit_numbers(str(out))
    assert all(f["check"] != "amount_words_mismatch" for f in res["findings"])


def test_date_format_inconsistent(tmp_path):
    doc = Document()
    doc.add_paragraph("起始 2024/01/01，結束 2024-12-31。")
    out = tmp_path / "date.docx"
    doc.save(str(out))
    res = na.audit_numbers(str(out))
    assert any(f["check"] == "date_format_inconsistent" for f in res["findings"])


def test_skill_cli_exit_codes(tmp_path):
    bad = _doc_with_table(tmp_path, [
        ["項目", "金額"], ["A", "1"], ["B", "1"], ["合計", "9"],
    ], name="bad")
    good = _doc_with_table(tmp_path, [
        ["項目", "金額"], ["A", "1"], ["B", "1"], ["合計", "2"],
    ], name="good")
    assert nau.main([good]) == 0
    assert nau.main([bad]) == 1
    assert nau.main([bad, "--fail-on", "high"]) == 1
