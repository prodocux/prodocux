from __future__ import annotations

import base64
import json
from copy import deepcopy
from importlib.resources import files
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from jsonschema import Draft202012Validator
from pptx import Presentation

from api.main import app
from prodocux_kernel.intake import (
    ProjectionValidationError,
    extract_pptx_continuable_projection,
    validate_continuable_projection,
)


def _presentation(slides: int = 55) -> bytes:
    presentation = Presentation()
    for number in range(1, slides + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"slide-{number}"
        slide.placeholders[1].text = (
            "TAIL_SLIDE_MARKER" if number == slides else f"body-{number}"
        )
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _schema() -> dict:
    return json.loads(
        files("prodocux_kernel.schemas")
        .joinpath("prodocux_pptx_continuable_projection_v1.json")
        .read_text(encoding="utf-8")
    )


def test_pptx_55_slides_continue_as_50_plus_5_with_shape_identity() -> None:
    raw = _presentation()
    first = extract_pptx_continuable_projection(raw)
    second = extract_pptx_continuable_projection(raw, cursor=first["next_cursor"])
    Draft202012Validator(_schema()).validate(first)
    Draft202012Validator(_schema()).validate(second)
    assert first["range"]["end"] == 51
    assert second["range"] == {
        "unit": "slide",
        "start": 51,
        "end": 56,
        "requested_max_slides": 50,
    }
    assert second["slides"][-1]["shapes"][1]["text"] == "TAIL_SLIDE_MARKER"
    assert second["slides"][-1]["shapes"][1]["shape_id"] > 0
    assert second["coverage"]["disposition"] == "complete"


def test_pptx_cursor_source_binding_fails_closed() -> None:
    raw = _presentation(2)
    first = extract_pptx_continuable_projection(raw, max_slides=1)
    cursor = deepcopy(first["next_cursor"])
    cursor["source_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="source_sha256"):
        extract_pptx_continuable_projection(raw, cursor=cursor, max_slides=1)


def test_pptx_continuation_http_and_capability() -> None:
    client = TestClient(app)
    request = {
        "document_filename": "large.pptx",
        "document_b64": base64.b64encode(_presentation()).decode(),
        "max_slides": 50,
    }
    first = client.post("/v1/intake/profile-presentation/continue", json=request)
    assert first.status_code == 200
    profiles = {
        item["format"]: item
        for item in client.get("/v1/intake/projection-capabilities").json()["profiles"]
    }
    assert profiles["pptx"]["range_units"] == ["slide"]


def test_pptx_semantics_reject_slide_identity_drift() -> None:
    result = extract_pptx_continuable_projection(_presentation(2), max_slides=1)
    invalid = deepcopy(result)
    invalid["slides"][0]["slide_number"] = 2
    with pytest.raises(ProjectionValidationError, match="contiguous"):
        validate_continuable_projection(invalid)
