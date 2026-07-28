"""Flagship pipeline tests: #2 profile extraction / #1 mapping / #3 rendering (including the #6 invariant gate)."""
from pathlib import Path

import pytest
from docx import Document

from prodocux_kernel.docops import fields as docfields
from prodocux_kernel.templating import profile, render
from prodocux_kernel.templating.mapping import (
    MappingConfig,
    evidence_pages,
    map_sources_to_sections,
    validate_config,
)
from prodocux_kernel.templating.render import FooterSpec, TocSpec


def _config():
    data = {
        "template_rules": {
            "heading_styles": ["Heading 1"],
            "body_style": "Normal",
            "remove_review_artifacts": ["{{", "}}"],
            "pagination": {"mode": "none"},
        },
        "sections": [
            {"id": "desc", "heading": "產品敘述", "mode": "paragraphs",
             "source_queries": ["product", "描述"]},
            {"id": "formula", "heading": "成分表", "mode": "table",
             "source_queries": ["INCI", "CAS"]},
        ],
    }
    return MappingConfig.from_dict(data)


def _template(tmp_path):
    doc = Document()
    doc.add_paragraph("目錄", style="Heading 1")
    doc.add_paragraph("舊目錄行 1 ........ 3")
    doc.add_paragraph("產品敘述", style="Heading 1")
    doc.add_paragraph("(待填)")
    doc.add_paragraph("成分表", style="Heading 1")
    doc.add_table(rows=2, cols=2)
    path = tmp_path / "template.docx"
    doc.save(str(path))
    return str(path)


# ---------------- #1 config schema ----------------
def test_config_from_dict():
    config = _config()
    assert config.heading_styles == ["Heading 1"]
    assert config.body_style == "Normal"
    assert len(config.sections) == 2
    assert config.sections[1].mode == "table"
    assert validate_config(config) == []


def test_pagination_policy_from_mapping_file():
    config = MappingConfig.from_json(
        Path(__file__).resolve().parents[1] / "examples" / "pif_tw" / "template_mapping.json"
    )
    assert config.pagination.mode == "mapped_sections_only"
    assert config.pagination.chapter_break_styles == ["PIF標題一"]
    assert config.table_policy == "reuse_template_table_shape_when_possible"
    assert len(config.front_matter) == 5
    basic = next(s for s in config.sections if s.id == "01_basic")
    assert basic.table_fill == "preserve_row_labels"


def test_config_validation_catches_bad_pagination_mode():
    data = {
        "template_rules": {
            "heading_styles": ["H1"],
            "body_style": "Normal",
            "pagination": {"mode": "invalid_mode"},
        },
        "sections": [{"id": "a", "heading": "H", "mode": "paragraphs"}],
    }
    issues = validate_config(MappingConfig.from_dict(data))
    assert any("pagination.mode" in i for i in issues)


def test_config_validation_catches_problems():
    data = {
        "template_rules": {"heading_styles": [], "body_style": "Normal"},
        "sections": [
            {"id": "a", "heading": "H", "mode": "paragraphs"},
            {"id": "a", "heading": "", "mode": "weird"},
        ],
    }
    issues = validate_config(MappingConfig.from_dict(data))
    assert any("heading_styles" in i for i in issues)
    assert any("重複" in i for i in issues)
    assert any("mode" in i for i in issues)


def test_headings_index():
    config = _config()
    idx = config.headings_index()
    assert set(idx.values()) == {"desc", "formula"}


# ---------------- #1 source mapping engine ----------------
def test_map_sources_to_sections():
    config = _config()
    pages = [
        {"file": "src.pdf", "page": 1, "text": "This product 描述 is nice"},
        {"file": "src.pdf", "page": 2, "text": "INCI Aqua CAS 7732-18-5 formula"},
        {"file": "src.pdf", "page": 3, "text": "irrelevant"},
    ]
    hits = map_sources_to_sections(pages, config)
    assert hits["desc"][0].page == 1
    assert hits["desc"][0].score == 2  # product + 描述 (description)
    assert hits["formula"][0].page == 2
    assert evidence_pages(hits, "formula") == "p.2"
    assert evidence_pages(hits, "missing_id") is None


# ---------------- #2 template profile extraction ----------------
def test_extract_profile(tmp_path):
    config = _config()
    tpl = _template(tmp_path)
    prof = profile.extract_profile(
        tpl, config.section_specs(), config.heading_styles, config.body_style
    )
    assert prof.missing == []
    desc = prof.section("desc")
    assert desc.found and desc.match == "exact"
    formula = prof.section("formula")
    assert formula.found and formula.has_table
    assert (formula.table_rows, formula.table_cols) == (2, 2)
    assert len(prof.discovered_headings) == 3


def test_extract_profile_missing(tmp_path):
    tpl = _template(tmp_path)
    specs = [{"id": "ghost", "heading": "不存在的標題", "mode": "paragraphs"}]
    prof = profile.extract_profile(tpl, specs, ["Heading 1"], "Normal")
    assert prof.missing == ["ghost"]
    assert prof.section("ghost").found is False


# ---------------- #3 render + #6 gate (end-to-end) ----------------
def test_render_document_passes_gate(tmp_path):
    config = _config()
    tpl = _template(tmp_path)
    out = tmp_path / "out.docx"
    drafts = {
        "desc": ["本產品為測試用香水。", "符合相關法規。"],
        "formula": ["成分內容（語言中立草稿）"],
    }
    result = render.render_document(
        tpl, out, drafts, config,
        toc=TocSpec(after_heading="目錄", until_heading="產品敘述"),
        footer=FooterSpec(prefix="第 ", middle=" 頁 / 共 ", suffix=" 頁"),
        review_terms=["待補充"],
    )
    assert result.passed is True
    assert result.sections_written == 2
    assert result.toc_inserted is True

    doc = Document(result.output_path)
    texts = [p.text for p in doc.paragraphs]
    assert "本產品為測試用香水。" in texts
    assert "(待填)" not in texts                 # old content has been replaced
    assert docfields.has_toc_field(doc) is True
    footer_xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml


def test_render_validation_can_be_skipped(tmp_path):
    config = _config()
    tpl = _template(tmp_path)
    out = tmp_path / "out2.docx"
    result = render.render_document(
        tpl, out, {"desc": ["x"]}, config, validate=False
    )
    assert result.invariants == []
    assert result.passed is True
