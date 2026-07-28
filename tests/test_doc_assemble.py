"""Skill doc_assemble tests: pipeline wrapping + bilingual report (i18n verification)."""
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from prodocux_kernel.templating.mapping import MappingConfig
from skills.doc_assemble import assemble as asm

CONFIG_PATH = Path(__file__).resolve().parents[1] / "examples" / "pif_tw" / "template_mapping.json"


def _config():
    return MappingConfig.from_json(CONFIG_PATH)


def _template(tmp_path, config):
    doc = Document()
    for name in (*config.heading_styles, config.body_style):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    h1, h2 = config.heading_styles
    doc.add_paragraph("目錄", style=h1)
    doc.add_paragraph("舊目錄 1 ... 3", style=config.body_style)
    doc.add_paragraph("產品敘述", style=h1)
    for spec in config.sections:
        doc.add_paragraph(spec.heading, style=h2)
        doc.add_paragraph("(待填)", style=config.body_style)
        if spec.mode == "table":
            doc.add_table(rows=2, cols=2)
    path = tmp_path / "tpl.docx"
    doc.save(str(path))
    return str(path)


def test_assemble_passes_gate(tmp_path):
    config = _config()
    tpl = _template(tmp_path, config)
    out = tmp_path / "out.docx"
    result = asm.assemble(
        tpl, config, {"02_registration": ["草稿一", "草稿二"]}, str(out),
        toc_after="目錄", toc_until="產品敘述", footer_style="zh",
        review_terms=["待補充"],
    )
    assert result.passed is True
    assert result.sections_written == 1
    assert result.toc_inserted is True


def test_report_is_bilingual(tmp_path):
    config = _config()
    tpl = _template(tmp_path, config)
    out = tmp_path / "out.docx"
    result = asm.assemble(tpl, config, {"02_registration": ["x"]}, str(out))

    en = asm.format_report(result, "en")
    zh = asm.format_report(result, "zh-TW")

    assert "Operations" in en and "Sections written" in en and "PASS" in en
    assert "處理動作" in zh and "寫入區段" in zh and "通過" in zh
    # invariant codes should also be localized (reuses #6's catalog)
    assert "valid_docx" in en and "可正常開啟" in zh


def test_cli_json_exit_code(tmp_path, capsys):
    config = _config()
    tpl = _template(tmp_path, config)
    out = tmp_path / "out.docx"
    drafts = tmp_path / "drafts.json"
    drafts.write_text(json.dumps({"11_stability": ["安定性草稿"]}, ensure_ascii=False),
                      encoding="utf-8")
    rc = asm.main([
        tpl, "--config", str(CONFIG_PATH), "--drafts", str(drafts),
        "--output", str(out), "--toc-after", "目錄", "--toc-until", "產品敘述",
        "--json",
    ])
    assert rc == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["passed"] is True
    assert payload["sections_written"] == 1
