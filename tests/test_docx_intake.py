from __future__ import annotations

import base64
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from api.main import app
from prodocux_kernel.intake import DOCX_PROFILE_SCHEMA, profile_docx_bytes


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Production Brief", level=1)
    document.add_paragraph("Unit A prepares the exterior scene.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Director"
    table.cell(1, 1).text = "Alice"
    document.sections[0].header.paragraphs[0].text = "Production Control"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_profile_docx_extracts_structure_without_interpretation() -> None:
    profile = profile_docx_bytes(_docx_bytes(), filename="brief.docx")
    assert profile["schema_version"] == DOCX_PROFILE_SCHEMA
    assert profile["heading_count"] == 1
    assert profile["table_count"] == 1
    assert profile["tables"][0]["preview"][1] == ["Director", "Alice"]
    assert profile["headers"] == ["Production Control"]
    assert profile["interpretation"] == "none"


def test_profile_document_api_accepts_small_base64() -> None:
    response = TestClient(app).post(
        "/v1/intake/profile-document",
        json={
            "document_b64": base64.b64encode(_docx_bytes()).decode("ascii"),
            "document_filename": "brief.docx",
        },
    )
    assert response.status_code == 200
    assert response.json()["profile"]["paragraph_count"] == 2


def test_invalid_docx_returns_client_error() -> None:
    response = TestClient(app).post(
        "/v1/intake/profile-document",
        json={
            "document_b64": base64.b64encode(b"not-docx").decode("ascii"),
            "document_filename": "broken.docx",
        },
    )
    assert response.status_code == 400
    assert response.json()["detail"] == "invalid DOCX document"
