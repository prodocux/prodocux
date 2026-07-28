"""#15 Document version diff — engine and skill tests."""
from docx import Document

from prodocux_kernel.docdiff import compare as dc
from skills.version_diff import version_diff as vd


def _doc(tmp_path, paragraphs, name):
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    out = tmp_path / f"{name}.docx"
    d.save(str(out))
    return str(out)


def test_no_change(tmp_path):
    a = _doc(tmp_path, ["第一段", "第二段"], "a")
    b = _doc(tmp_path, ["第一段", "第二段"], "b")
    res = dc.compare_docx(a, b)
    assert res["changed"] is False
    assert res["summary"] == {"added": 0, "removed": 0, "replaced": 0}


def test_added_removed_replaced(tmp_path):
    a = _doc(tmp_path, ["保留", "舊內容", "待刪"], "old")
    b = _doc(tmp_path, ["保留", "新內容", "新增段"], "new")
    res = dc.compare_docx(a, b)
    assert res["changed"] is True
    types = [c["type"] for c in res["changes"]]
    assert "replaced" in types
    # "待刪" removed, "新增段" added (may be merged into a replace, depending on the opcode)
    assert res["summary"]["replaced"] >= 1


def test_table_change(tmp_path):
    d1 = Document()
    t1 = d1.add_table(rows=1, cols=2)
    t1.cell(0, 0).text = "價格"
    t1.cell(0, 1).text = "100"
    p1 = tmp_path / "t1.docx"
    d1.save(str(p1))

    d2 = Document()
    t2 = d2.add_table(rows=1, cols=2)
    t2.cell(0, 0).text = "價格"
    t2.cell(0, 1).text = "200"
    p2 = tmp_path / "t2.docx"
    d2.save(str(p2))

    res = dc.compare_docx(str(p1), str(p2))
    assert res["changed"] is True


def test_skill_cli(tmp_path):
    a = _doc(tmp_path, ["同樣"], "same1")
    b = _doc(tmp_path, ["同樣"], "same2")
    assert vd.main([a, b]) == 0
    assert vd.main([a, b, "--fail-on-change"]) == 0

    c = _doc(tmp_path, ["改了"], "changed")
    assert vd.main([a, c]) == 0  # does not fail on differences by default
    assert vd.main([a, c, "--fail-on-change"]) == 1
