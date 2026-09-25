from __future__ import annotations

import base64
import copy
import io
import json
import zipfile
from importlib import resources

import pytest
from docx import Document
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient
from jsonschema import Draft202012Validator

from api.main import app
from prodocux_kernel.rendering import (
    extract_content_blocks,
    extract_continuable_projection,
)
from prodocux_kernel.rendering.errors import RenderContractError
from prodocux_kernel.rendering.validate import validate_continuable_projection_result


def test_projection_schema_is_valid_draft_2020_12() -> None:
    projection_schema = json.loads(
        (
            resources.files("prodocux_kernel.schemas")
            / "prodocux_continuable_projection_v1.json"
        ).read_text(encoding="utf-8")
    )
    cursor_schema = json.loads(
        (
            resources.files("prodocux_kernel.schemas")
            / "prodocux_projection_cursor_v1.json"
        ).read_text(encoding="utf-8")
    )
    Draft202012Validator.check_schema(projection_schema)
    Draft202012Validator.check_schema(cursor_schema)
    embedded = projection_schema["$defs"]["cursor"]
    assert embedded["description"] == cursor_schema["description"]
    assert embedded["type"] == cursor_schema["type"]
    assert embedded["additionalProperties"] == cursor_schema["additionalProperties"]
    assert embedded["required"] == cursor_schema["required"]
    assert embedded["properties"] == cursor_schema["properties"]


def _heading_docx(count: int, *, tail: str = "TAIL-MARKER-206") -> bytes:
    document = Document()
    for index in range(1, count + 1):
        text = tail if index == count else f"Heading {index}"
        document.add_heading(text, level=1)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _docx_with_out_of_scope_content() -> bytes:
    document = Document()
    for index in range(1, 207):
        document.add_heading(f"Body {index}", level=1)
    document.sections[0].header.paragraphs[0].text = "HEADER_UNIQUE_MARKER"
    document.sections[0].footer.paragraphs[0].text = "FOOTER_UNIQUE_MARKER"
    body = document._element.body
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "UNSUPPORTED_BODY_MARKER"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    control.append(content)
    body.insert(len(body) - 1, control)
    stream = io.BytesIO()
    document.save(stream)

    source = io.BytesIO(stream.getvalue())
    target = io.BytesIO()
    with (
        zipfile.ZipFile(source) as current,
        zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as updated,
    ):
        for item in current.infolist():
            updated.writestr(item, current.read(item.filename))
        updated.writestr(
            "word/comments.xml",
            '<w:comments xmlns:w="http://schemas.openxmlformats.org/'
            'wordprocessingml/2006/main"><w:comment w:id="0"><w:p><w:r>'
            "<w:t>COMMENT_UNIQUE_MARKER</w:t></w:r></w:p></w:comment></w:comments>",
        )
        updated.writestr(
            "word/footnotes.xml",
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/'
            'wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r>'
            "<w:t>FOOTNOTE_UNIQUE_MARKER</w:t></w:r></w:p></w:footnote></w:footnotes>",
        )
    return target.getvalue()


def test_legacy_contract_stays_bounded_while_projection_continues() -> None:
    payload = _heading_docx(206)

    legacy = extract_content_blocks("large.docx", payload)
    assert len(legacy["content"]["blocks"]) == 200
    assert legacy["truncated"] is True

    first = extract_continuable_projection("large.docx", payload)
    assert first["range"] == {
        "unit": "block",
        "start": 0,
        "end_exclusive": 200,
        "returned_blocks": 200,
        "requested_max_blocks": 200,
    }
    assert first["coverage"]["disposition"] == "partial_unknown"
    assert first["coverage"]["continuation_available"] is True
    assert first["coverage"]["known_total_blocks"] is None
    assert first["counts"]["processed_through_range_end"]["blocks"] == 200
    assert first["counts"]["known_total"]["blocks"] is None
    assert first["counts"]["processed_through_range_end"]["pages"] is None
    assert first["ocr_disposition"] == "not_applicable"

    second = extract_continuable_projection(
        "large.docx", payload, cursor=first["next_cursor"]
    )
    assert second["range"]["start"] == 200
    assert second["range"]["end_exclusive"] == 206
    assert second["coverage"]["disposition"] == "complete"
    assert second["coverage"]["known_total_blocks"] == 206
    assert second["counts"]["known_total"]["blocks"] == 206
    assert (
        second["counts"]["known_total"]["text_utf8_bytes"]
        == second["counts"]["processed_through_range_end"]["text_utf8_bytes"]
    )
    assert second["next_cursor"] is None
    assert second["content"]["blocks"][-1]["text"] == "TAIL-MARKER-206"

    combined = first["content"]["blocks"] + second["content"]["blocks"]
    assert [block["id"] for block in combined] == [f"h{i}" for i in range(1, 207)]
    expected_text_bytes = sum(len(block["text"].encode("utf-8")) for block in combined)
    assert second["counts"]["processed_through_range_end"]["text_utf8_bytes"] == (
        expected_text_bytes
    )
    assert (
        second["counts"]["processed_through_range_end"]["text_utf8_bytes"]
        > (first["counts"]["processed_through_range_end"]["text_utf8_bytes"])
    )


def test_projection_is_deterministic_and_source_bound() -> None:
    payload = _heading_docx(4)
    first = extract_continuable_projection("small.docx", payload, max_blocks=2)
    assert first == extract_continuable_projection("small.docx", payload, max_blocks=2)

    mutated = _heading_docx(5)
    with pytest.raises(RenderContractError) as source_error:
        extract_continuable_projection(
            "small.docx", mutated, cursor=first["next_cursor"], max_blocks=2
        )
    assert source_error.value.code == "CONTINUATION_SOURCE_MISMATCH"

    caller_selected = dict(first["next_cursor"])
    caller_selected["next_block"] = 3
    selected = extract_continuable_projection(
        "small.docx", payload, cursor=caller_selected, max_blocks=2
    )
    assert selected["range"]["start"] == 3
    assert selected["range"]["start"] != first["range"]["end_exclusive"]


def test_out_of_scope_docx_content_is_disclosed_before_first_page() -> None:
    result = extract_continuable_projection(
        "scope.docx", _docx_with_out_of_scope_content(), max_blocks=1
    )
    assert result["coverage"]["disposition"] == "partial_unknown"
    omitted = set(result["coverage"]["omitted_content_classes"])
    assert {
        "docx_headers",
        "docx_footers",
        "docx_comments",
        "docx_footnotes",
        "unsupported_docx_body_elements",
    } <= omitted
    projected = json.dumps(result["content"])
    assert "HEADER_UNIQUE_MARKER" not in projected
    assert "UNSUPPORTED_BODY_MARKER" not in projected


def test_projection_semantics_reject_inconsistent_derived_fields() -> None:
    payload = _heading_docx(3)
    result = extract_continuable_projection("large.docx", payload, max_blocks=2)
    wrong_text = copy.deepcopy(result)
    wrong_text["text_items"][0]["text"] = "not derived from content"
    with pytest.raises(RenderContractError):
        validate_continuable_projection_result(wrong_text)

    wrong_binding = copy.deepcopy(result)
    wrong_binding["next_cursor"]["source_sha256"] = "0" * 64
    with pytest.raises(RenderContractError):
        validate_continuable_projection_result(wrong_binding)

    wrong_coverage = copy.deepcopy(result)
    wrong_coverage["coverage"]["disposition"] = "partial_known"
    with pytest.raises(RenderContractError):
        validate_continuable_projection_result(wrong_coverage)

    exceeds_requested = copy.deepcopy(result)
    exceeds_requested["range"]["requested_max_blocks"] = 1
    with pytest.raises(RenderContractError):
        validate_continuable_projection_result(exceeds_requested)

    final = extract_continuable_projection("final.docx", _heading_docx(1))
    invalid_combinations = []

    nonfinal_with_total = copy.deepcopy(result)
    nonfinal_with_total["coverage"]["known_total_blocks"] = 3
    nonfinal_with_total["counts"]["known_total"]["blocks"] = 3
    invalid_combinations.append(nonfinal_with_total)

    final_partial_unknown = copy.deepcopy(final)
    final_partial_unknown["coverage"]["disposition"] = "partial_unknown"
    invalid_combinations.append(final_partial_unknown)

    final_complete_with_omission = copy.deepcopy(final)
    final_complete_with_omission["coverage"]["omitted_content_classes"] = [
        "unprojected"
    ]
    invalid_combinations.append(final_complete_with_omission)

    flag_without_cursor = copy.deepcopy(result)
    flag_without_cursor["next_cursor"] = None
    invalid_combinations.append(flag_without_cursor)

    for invalid in invalid_combinations:
        with pytest.raises(RenderContractError):
            validate_continuable_projection_result(invalid)

    second = extract_continuable_projection(
        "large.docx",
        payload,
        cursor=result["next_cursor"],
        max_blocks=2,
    )
    wrong_request_cursor = dict(result["next_cursor"])
    wrong_request_cursor["parser_contract_name"] = "other_parser_v1"
    with pytest.raises(RenderContractError):
        validate_continuable_projection_result(
            second, request_cursor=wrong_request_cursor
        )


def test_unquantified_text_and_table_clipping_is_partial_unknown() -> None:
    text_document = Document()
    text_document.add_paragraph("x" * 9000)
    text_stream = io.BytesIO()
    text_document.save(text_stream)
    text_result = extract_continuable_projection("text.docx", text_stream.getvalue())
    assert text_result["coverage"]["disposition"] == "partial_unknown"
    assert (
        "paragraph_text_beyond_limit"
        in text_result["coverage"]["omitted_content_classes"]
    )

    table_document = Document()
    table_document.add_table(rows=501, cols=1)
    table_document.add_table(rows=1, cols=33)
    table_stream = io.BytesIO()
    table_document.save(table_stream)
    table_result = extract_continuable_projection(
        "tables.docx", table_stream.getvalue()
    )
    assert table_result["coverage"]["disposition"] == "partial_unknown"
    assert {
        "table_rows_beyond_limit",
        "table_columns_beyond_limit",
    } <= set(table_result["coverage"]["omitted_content_classes"])


def test_http_projection_round_trip_and_format_boundary() -> None:
    payload = _heading_docx(3)
    client = TestClient(app)
    capabilities = client.get("/v1/intake/capabilities").json()
    docx = next(
        item for item in capabilities["formats"] if ".docx" in item["extensions"]
    )
    assert "extract_blocks_continue" in docx["additional_operations"]
    request = {
        "document_filename": "sample.docx",
        "document_b64": base64.b64encode(payload).decode("ascii"),
        "max_blocks": 2,
    }
    first = client.post("/v1/intake/extract-blocks/continue", json=request)
    assert first.status_code == 200
    cursor = first.json()["next_cursor"]

    request["cursor"] = cursor
    second = client.post("/v1/intake/extract-blocks/continue", json=request)
    assert second.status_code == 200
    assert second.json()["range"]["start"] == 2

    request["document_b64"] = base64.b64encode(_heading_docx(4)).decode("ascii")
    changed = client.post("/v1/intake/extract-blocks/continue", json=request)
    assert changed.status_code == 409
    assert changed.json()["detail"]["code"] == "CONTINUATION_SOURCE_MISMATCH"

    unsupported = client.post(
        "/v1/intake/extract-blocks/continue",
        json={
            "document_filename": "sample.pdf",
            "document_b64": base64.b64encode(b"%PDF-1.4").decode("ascii"),
        },
    )
    assert unsupported.status_code == 501
