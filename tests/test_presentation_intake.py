from __future__ import annotations

import base64
from io import BytesIO

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from api.main import app
from prodocux_kernel.intake import PRESENTATION_PROFILE_SCHEMA, profile_pptx_bytes


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Production Plan"
    slide.placeholders[1].text = "Unit A exterior\nUnit B interior"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Unit"
    table.cell(0, 1).text = "Call"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "06:30"
    slide.notes_slide.notes_text_frame.text = "Confirm location permit."
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_profile_pptx_extracts_slide_content_without_interpretation() -> None:
    profile = profile_pptx_bytes(_pptx_bytes(), filename="plan.pptx")
    assert profile["schema_version"] == PRESENTATION_PROFILE_SCHEMA
    assert profile["slide_count"] == 1
    assert profile["interpretation"] == "none"
    slide = profile["slides"][0]
    assert slide["title"] == "Production Plan"
    assert "Unit A exterior" in slide["texts"][0]
    assert slide["speaker_notes"] == "Confirm location permit."
    assert slide["tables"][0]["preview"][1] == ["A", "06:30"]
    assert slide["shape_count_is_lower_bound"] is False


def test_profile_presentation_api_accepts_small_base64() -> None:
    response = TestClient(app).post(
        "/v1/intake/profile-presentation",
        json={
            "document_b64": base64.b64encode(_pptx_bytes()).decode("ascii"),
            "document_filename": "plan.pptx",
        },
    )
    assert response.status_code == 200
    assert response.json()["profile"]["slides"][0]["title"] == "Production Plan"


def test_invalid_pptx_returns_client_error() -> None:
    response = TestClient(app).post(
        "/v1/intake/profile-presentation",
        json={
            "document_b64": base64.b64encode(b"not-pptx").decode("ascii"),
            "document_filename": "broken.pptx",
        },
    )
    assert response.status_code == 400
    assert response.json()["detail"] == "invalid PPTX presentation"
