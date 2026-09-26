from __future__ import annotations

import base64
import json
from copy import deepcopy
from importlib.resources import files
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from jsonschema import Draft202012Validator
from openpyxl import Workbook

from api.main import app
from prodocux_kernel.intake import (
    ProjectionValidationError,
    extract_xlsx_continuable_projection,
    validate_continuable_projection,
)


def _workbook() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Primary"
    for row in range(1, 507):
        first.cell(row=row, column=1, value=f"row-{row}")
        first.cell(row=row, column=2, value=f"=A{row}")
    second = workbook.create_sheet("Tail")
    second["A1"] = "TAIL_SHEET_MARKER"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _schema() -> dict:
    return json.loads(
        files("prodocux_kernel.schemas")
        .joinpath("prodocux_xlsx_continuable_projection_v1.json")
        .read_text(encoding="utf-8")
    )


def test_xlsx_continues_rows_then_advances_to_next_sheet() -> None:
    raw = _workbook()
    first = extract_xlsx_continuable_projection(raw)
    second = extract_xlsx_continuable_projection(raw, cursor=first["next_cursor"])
    third = extract_xlsx_continuable_projection(raw, cursor=second["next_cursor"])
    for result in (first, second, third):
        Draft202012Validator(_schema()).validate(result)
    assert (
        first["range"]["sheet_name"],
        first["range"]["start"],
        first["range"]["end"],
    ) == ("Primary", 1, 501)
    assert second["range"]["start"] == 501 and second["range"]["end"] == 507
    assert third["range"]["sheet_name"] == "Tail"
    assert third["rows"] == [["TAIL_SHEET_MARKER"]]
    assert third["coverage"]["disposition"] == "complete"
    assert first["formula_cells"][0]["cell"] == "B1"


def test_xlsx_cursor_binds_source_and_sheet_name() -> None:
    raw = _workbook()
    first = extract_xlsx_continuable_projection(raw)
    wrong = deepcopy(first["next_cursor"])
    wrong["next_sheet_name"] = "forged"
    with pytest.raises(ValueError, match="sheet name"):
        extract_xlsx_continuable_projection(raw, cursor=wrong)


def test_xlsx_continuation_http_and_capability() -> None:
    client = TestClient(app)
    request = {
        "document_filename": "large.xlsx",
        "document_b64": base64.b64encode(_workbook()).decode(),
        "max_rows": 500,
    }
    first = client.post("/v1/intake/profile-workbook/continue", json=request)
    assert first.status_code == 200
    profiles = {
        item["format"]: item
        for item in client.get("/v1/intake/projection-capabilities").json()["profiles"]
    }
    assert profiles["xlsx"]["range_units"] == ["worksheet_row"]


def test_xlsx_semantics_reject_same_sheet_gap() -> None:
    result = extract_xlsx_continuable_projection(_workbook())
    invalid = deepcopy(result)
    invalid["next_cursor"]["next_row"] += 1
    with pytest.raises(ProjectionValidationError, match="same-sheet"):
        validate_continuable_projection(invalid)
