"""Table-section writing tests (the ingredient table §03 must not be skipped)."""
import json
import os
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from prodocux_kernel.audit.pif_tw import audit_pif_tw, extract_section_bodies, load_checklist
from prodocux_kernel.docops.content import parse_draft_entry, parse_drafts
from prodocux_kernel.docops.tables import fill_table_data, set_cell_text
from prodocux_kernel.templating.mapping import MappingConfig
from prodocux_kernel.templating import render
from prodocux_kernel.templating.profile import extract_profile

EXAMPLES = Path(__file__).resolve().parents[1] / "examples" / "pif_tw"
_REAL_TPL = os.environ.get("PRODOCUX_TW_PIF_TEMPLATE", "").strip()
REAL_TPL = Path(_REAL_TPL) if _REAL_TPL else Path("__unset_PRODOCUX_TW_PIF_TEMPLATE__")

FORMULA_ROWS = [
    ["Aqua", "70.00000", "7732-18-5", "溶劑"],
    ["Alcohol denat.", "25.00000", "64-17-5", "溶劑"],
    ["Parfum", "5.00000", "-", "香精"],
]


def _mapping():
    return MappingConfig.from_json(EXAMPLES / "template_mapping.json")


def _formula_template(tmp_path):
    config = _mapping()
    doc = Document()
    for name in (*config.heading_styles, config.body_style):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    h2 = config.heading_styles[1]
    doc.add_paragraph("全成分名稱及其各別含量", style=h2)
    t = doc.add_table(rows=2, cols=4)
    set_cell_text(t.cell(0, 0), "INCI 名稱")
    set_cell_text(t.cell(0, 1), "含量 (%)")
    set_cell_text(t.cell(0, 2), "CAS No.")
    set_cell_text(t.cell(0, 3), "用途")
    path = tmp_path / "formula_tpl.docx"
    doc.save(str(path))
    return str(path), config


def test_parse_draft_entry_table():
    entry = {"table": {"header_rows": 1, "rows": FORMULA_ROWS}}
    content = parse_draft_entry(entry)
    assert content.has_table
    assert len(content.table_rows) == 3
    assert content.table_header_rows == 1


def test_parse_draft_entry_by_label():
    entry = {
        "table": {
            "header_rows": 1,
            "fill_mode": "by_label",
            "values": {"產品名稱(中文/英文)": "測試品名", "產品類別": "香水"},
        }
    }
    content = parse_draft_entry(entry)
    assert content.has_table
    assert content.table_fill_mode == "by_label"
    assert content.table_values["產品類別"] == "香水"


def test_parse_drafts_mixed():
    raw = {
        "02_registration": ["CPNP 草稿"],
        "03_formula": {"table": {"header_rows": 1, "rows": FORMULA_ROWS}},
    }
    parsed = parse_drafts(raw)
    assert parsed["02_registration"].has_paragraphs
    assert parsed["03_formula"].has_table


def test_write_formula_table_synthetic(tmp_path):
    tpl, config = _formula_template(tmp_path)
    out = tmp_path / "out.docx"
    drafts = {
        "03_formula": {
            "table": {"header_rows": 1, "rows": FORMULA_ROWS},
        }
    }
    result = render.render_document(tpl, out, drafts, config, validate=False)
    assert result.tables_written == 1

    prof = extract_profile(tpl, config.section_specs(), config.heading_styles, config.body_style)
    bodies = extract_section_bodies(Document(str(out)), prof)
    body = bodies["03_formula"]
    assert "INCI" in body or "Aqua" in body
    assert "Alcohol denat." in body
    assert "7732-18-5" in body

    checklist = load_checklist(EXAMPLES / "tw_pif_checklist.json")
    audit = audit_pif_tw(out, config, checklist)
    formula_fails = [f for f in audit["findings"] if f["section_id"] == "03_formula"]
    assert not any(f["check"] == "section_empty_or_placeholder" for f in formula_fails)
    assert not any(f["check"] == "section_required_keyword_missing" for f in formula_fails)


@pytest.mark.skipif(not REAL_TPL.is_file(), reason="Set PRODOCUX_TW_PIF_TEMPLATE for optional real-template write test")
def test_write_formula_table_real_template(tmp_path):
    config = _mapping()
    out = tmp_path / "real_out.docx"
    drafts = {"03_formula": {"table": {"header_rows": 1, "rows": FORMULA_ROWS}}}
    result = render.render_document(
        str(REAL_TPL), out, drafts, config,
        toc=render.TocSpec(after_heading="目錄", until_heading="產品敘述"),
        footer=render.FooterSpec(prefix="第 ", middle=" 頁 / 共 ", suffix=" 頁"),
    )
    assert result.tables_written == 1
    assert result.passed is True

    prof = extract_profile(str(out), config.section_specs(), config.heading_styles, config.body_style)
    bodies = extract_section_bodies(Document(str(out)), prof)
    assert "Aqua" in bodies.get("03_formula", "")
    assert "7732-18-5" in bodies.get("03_formula", "")

    checklist = load_checklist(EXAMPLES / "tw_pif_checklist.json")
    audit = audit_pif_tw(out, config, checklist)
    f03 = [f for f in audit["findings"] if f["section_id"] == "03_formula"]
    assert not any(f["check"] == "section_required_keyword_missing" for f in f03)


def test_fill_table_data_unit():
    doc = Document()
    t = doc.add_table(rows=3, cols=2)
    n = fill_table_data(t, [["a", "b"], ["c", "d"]], start_row=1)
    assert n == 4
    assert t.cell(1, 0).text == "a"
    assert t.cell(2, 1).text == "d"
