"""Flagship pipeline: intake → mapping → render (+ evidence image ops)."""
import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from prodocux_kernel.docops import tables
from prodocux_kernel.intake import pdf as intake
from prodocux_kernel.intake import pdf_images
from prodocux_kernel.templating.mapping import MappingConfig
from prodocux_kernel.templating import pipeline
from skills.doc_assemble import assemble as asm

CONFIG_PATH = Path(__file__).resolve().parents[1] / "examples" / "pif_tw" / "template_mapping.json"
PAGES_SAMPLE = Path(__file__).resolve().parents[1] / "examples" / "pif_tw" / "sample_pages.json"


def _mini_config():
    return MappingConfig.from_dict({
        "template_rules": {"heading_styles": ["H1"], "body_style": "Body"},
        "sections": [
            {"id": "02_registration", "heading": "完成產品登錄之證明文件", "mode": "paragraphs",
             "source_queries": ["CPNP", "notification"]},
        ],
    })


def _template(tmp_path, config):
    doc = Document()
    for name in (*config.heading_styles, config.body_style):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("目錄", style=config.heading_styles[0])
    for spec in config.sections:
        doc.add_paragraph(spec.heading, style=config.heading_styles[0])
        doc.add_paragraph("(待填)", style=config.body_style)
    path = tmp_path / "tpl.docx"
    doc.save(str(path))
    return str(path)


def test_pipeline_pages_to_source_map(tmp_path):
    if not PAGES_SAMPLE.is_file():
        pytest.skip("sample_pages.json missing")
    config = _mini_config()
    tpl = _template(tmp_path, config)
    out = tmp_path / "out.docx"
    smap = tmp_path / "source_map.json"
    pr = pipeline.run_flagship_pipeline(
        tpl, out, config, {"02_registration": ["CPNP 已通報"]},
        pages_json=str(PAGES_SAMPLE),
        source_map_out=str(smap),
        validate=False,
    )
    assert pr.source_map_path == str(smap)
    data = json.loads(smap.read_text(encoding="utf-8"))
    assert data["schema"] == pipeline.SOURCE_MAP_SCHEMA
    assert "02_registration" in data["sections"]
    assert pr.render.sections_written == 1


def test_doc_assemble_with_pages_json(tmp_path):
    if not PAGES_SAMPLE.is_file():
        pytest.skip("sample_pages.json missing")
    config = _mini_config()
    tpl = _template(tmp_path, config)
    out = tmp_path / "out.docx"
    result = asm.assemble(
        tpl, config, {"02_registration": ["草稿"]}, str(out),
        pages_json=str(PAGES_SAMPLE),
        source_map_out=str(tmp_path / "map.json"),
        validate=False,
    )
    pr = asm.last_pipeline_result()
    assert pr is not None
    assert pr.pages_document["page_count"] > 0
    assert result.sections_written == 1


def test_fill_image_in_table_row_by_label(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "內包裝/容器"
    table.rows[0].cells[1].text = ""
    img = tmp_path / "x.png"
    pymupdf = pytest.importorskip("fitz")
    writer = pymupdf.open()
    page = writer.new_page()
    page.insert_text((50, 50), "img")
    writer.save(str(tmp_path / "src.pdf"))
    writer.close()
    png, err = pdf_images.render_pdf_page(tmp_path / "src.pdf", 1, img)
    assert err is None and png
    ok = tables.fill_image_in_table_row_by_label(doc, "內包裝/容器", png, caption="照片")
    assert ok is True
    out = tmp_path / "filled.docx"
    doc.save(str(out))
    doc2 = Document(str(out))
    assert "照片" in doc2.tables[0].rows[0].cells[1].text


def test_render_pdf_page_requires_pymupdf(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf = tmp_path / "blank.pdf"
    writer.write(str(pdf))
    pymupdf = pytest.importorskip("fitz")
    path, err = pdf_images.render_pdf_page(pdf, 1, tmp_path / "p1.png")
    assert err is None
    assert path and path.is_file()


def test_evidence_spec_schema():
    spec_path = Path(__file__).resolve().parents[1] / "examples" / "pif_tw" / "evidence_spec.json"
    from prodocux_kernel.intake.evidence import load_evidence_spec
    spec = load_evidence_spec(spec_path)
    assert len(spec["extract"]) >= 5
    assert any(i["type"] == "appendix" for i in spec["inject"])
