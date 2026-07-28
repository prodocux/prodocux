"""Skill #6 end-to-end tests (CLI / report layer)."""
from docx import Document
from docx.enum.text import WD_BREAK

from skills.structure_health import health_check as hc


def _clean(tmp_path):
    doc = Document()
    doc.add_paragraph("一般內文")
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("第二頁內容")
    out = tmp_path / "clean.docx"
    doc.save(str(out))
    return str(out)


def _broken(tmp_path):
    doc = Document()
    doc.add_paragraph("目錄")
    doc.add_paragraph("第一章 簡介 ............... 1")
    doc.add_paragraph("詳見 錯誤! 找不到參照來源。")
    out = tmp_path / "broken.docx"
    doc.save(str(out))
    return str(out)


def test_check_clean_passes(tmp_path):
    rep = hc.check_file(_clean(tmp_path))
    assert rep["passed"] is True
    assert rep["fail_count"] == 0


def test_check_broken_fails_with_severity(tmp_path):
    rep = hc.check_file(_broken(tmp_path))
    assert rep["passed"] is False
    # hardcoded TOC + broken reference => high
    assert rep["max_severity"] == "high"
    ids = {c["id"]: c for c in rep["checks"]}
    assert ids["toc_is_field"]["passed"] is False
    assert ids["cross_references_valid"]["passed"] is False
    # the kernel outputs a language-neutral code, not a Chinese sentence
    assert ids["toc_is_field"]["code"] == "toc_is_field.hardcoded"


def test_format_report_bilingual(tmp_path):
    rep = hc.check_file(_broken(tmp_path))
    en = hc.format_report(rep, "en")
    assert "FAIL" in en and "Fix:" in en
    zh = hc.format_report(rep, "zh-TW")
    assert "未通過" in zh and "建議：" in zh


def test_main_exit_codes(tmp_path):
    assert hc.main([_clean(tmp_path)]) == 0
    assert hc.main([_broken(tmp_path)]) == 1
    # fail-on critical: broken only has "high" severity, so it should not fail
    assert hc.main([_broken(tmp_path), "--fail-on", "critical"]) == 0
    assert hc.main([_broken(tmp_path), "--fail-on", "high"]) == 1
